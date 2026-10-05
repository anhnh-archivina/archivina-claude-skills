#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Soat mat bang tran can ho (RCP) theo bo quy tac Archivina, bao vi tri khong phu hop va DE XUAT vi tri moi.

Doc DXF (da xuat bang dwg_to_dxf_aec.ps1 cua skill dien-tich-ch): nhan dien phong (ten phong + net ranh), noi that
(tu ao, giuong, ban an, sofa, thiet bi ve sinh) va thiet bi tran (theo assets/thu-vien/catalog.json), roi kiem tra:
  - den chung: khoang cach tam-tam >= 1200; cach tuong >= 500 (uu tien ~600)          -> HARD-RULE WARNING
  - moi thiet bi trong vung tu ao                                                    -> CONFLICT (PCCC: CRITICAL)
  - thiet bi chong len nhau                                                           -> COORDINATION WARNING
  - phong ngu: den tren vung goi, gio cap thoi vao vung goi, lo tham tren giuong      -> DESIGN / COORDINATION
  - WC: den theo truc thiet bi ve sinh, den roi guong theo truc guong                 -> DESIGN WARNING
  - den tha theo tam ban an / bo sofa; lo tham giua phong khach; dau bao gan gio cap   -> DESIGN / COORDINATION
De xuat: bo tri lai den chung cua phong vi pham HARD-RULE (giam so den thay vi ep khoang cach), doi thiet bi
khong phai PCCC ra ngoai vung cam; thiet bi PCCC KHONG tu doi, chi bao TECHNICAL REVIEW REQUIRED.
Khong sua DXF dau vao. Xuat: JSON (stdout), BaoCaoSoatTran.xlsx, xem_tran_<can>.png, ve_de_xuat_tran.scr.

    python soat_tran.py <file.dxf> --out-dir <thu muc> [--du-an "Ten"] [--thu-vien <thu muc dwg block>]
"""
import argparse
import fnmatch
import importlib.util
import io
import json
import math
import os
import re
import sys
import time
import unicodedata
from collections import defaultdict

import ezdxf
import shapely
from ezdxf import bbox
from shapely.geometry import LineString, MultiPoint, Point, Polygon, box
from shapely.ops import nearest_points, unary_union

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
_tpp_dir = os.environ.get("DIEN_TICH_CH_SCRIPTS") or os.path.join(os.path.dirname(SKILL), "dien-tich-ch", "scripts")
_spec = importlib.util.spec_from_file_location("tpp", os.path.join(_tpp_dir, "tao_polyline_phong.py"))
tpp = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(tpp)

CRIT, HARD, COORD, DESIGN = "CRITICAL", "HARD-RULE WARNING", "COORDINATION WARNING", "DESIGN WARNING"
MUC = {CRIT: "Lỗi", HARD: "Lỗi", COORD: "Cảnh báo", DESIGN: "Gợi ý"}
PCCC = {"FIRE_ALARM", "SPRINKLER"}


def bo_dau(s):
    s = unicodedata.normalize("NFD", (s or "").replace("đ", "d").replace("Đ", "D"))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").lower().strip()


def ten_goc(name):
    """Ten block sau khi bo tien to xref da bind ('A$0$B$0$C' -> 'C') va bo dau."""
    return bo_dau(re.split(r"\$\d+\$", name or "")[-1])


def khop(name, patterns):
    n = ten_goc(name)
    return any(fnmatch.fnmatchcase(n, bo_dau(p)) for p in patterns)


def lop(e):
    return tpp.layer_goc(e.dxf.layer)


def phang(ents, depth=0):
    """Trai phang doi tuong long nhau (da bien doi toa do ve WCS)."""
    for e in ents:
        if e.dxftype() == "INSERT" and depth < 4:
            try:
                yield from phang(e.virtual_entities(), depth + 1)
            except Exception:
                continue
        else:
            yield e


def hinh_ky_hieu(ins):
    """Hop bao (WCS) cua phan ky hieu that (bo Defpoints, A-PCCC-Phu, vong tron phu R > 900)."""
    core = []
    for e in phang(ins.virtual_entities()):
        if lop(e) in ("defpoints", "a-pccc-phu"):
            continue
        if e.dxftype() == "CIRCLE" and e.dxf.radius > 900:
            continue
        core.append(e)
    if not core:
        return None
    try:
        ext = bbox.extents(core, fast=True)
    except Exception:
        return None
    if not ext.has_data:
        return None
    return box(ext.extmin.x, ext.extmin.y, ext.extmax.x, ext.extmax.y)


_BB = {}


def bb_ins(e, doc):
    """Hop bao (WCS) cua INSERT: tinh extents cua DINH NGHIA block mot lan (cache) roi bien doi 4 goc."""
    name = e.dxf.name
    if name not in _BB:
        try:
            ext = bbox.extents(doc.blocks.get(name), fast=True)
            _BB[name] = (ext.extmin, ext.extmax) if ext.has_data else None
        except Exception:
            _BB[name] = None
    v = _BB[name]
    if v is None:
        return None
    m = e.matrix44()
    (x0, y0), (x1, y1) = (v[0].x, v[0].y), (v[1].x, v[1].y)
    pts = [m.transform((x, y, 0)) for x, y in ((x0, y0), (x1, y0), (x0, y1), (x1, y1))]
    xs, ys = [p.x for p in pts], [p.y for p in pts]
    return box(min(xs), min(ys), max(xs), max(ys))


# ----------------------------------------------------------------------------------------------- nhan dien thiet bi
class NhanDien:
    def __init__(self, catalog, cfg):
        self.cat = catalog["thiet_bi"]
        self.lay_tb = {tpp.layer_goc(x) for x in cfg["layer_thiet_bi"]}

    def theo_ten(self, name):
        for c in self.cat:
            if khop(name, c["bi_danh"]):
                return c
        return None

    def theo_chu_ky(self, layer, w, h):
        for c in self.cat:
            k = c.get("chu_ky")
            if not k:
                continue
            if tpp.layer_goc(layer) in {tpp.layer_goc(x) for x in k["layer"]} and \
                    (k["rong"][0] <= w <= k["rong"][1] and k["cao"][0] <= h <= k["cao"][1]):
                return c
        return None


def doc_thiet_bi(msp, doc, nd, vung):
    """Tim block thiet bi tran trong vung (polygon) - di sau toi 3 cap. Tra ve (list thiet bi, list block chua biet)."""
    out, la = [], []

    def con_insert(name):
        b = doc.blocks.get(name)
        return [x for x in b if x.dxftype() == "INSERT"] if b else []

    def duyet(ents, depth):
        for e in ents:
            if e.dxftype() != "INSERT":
                continue
            bb = bb_ins(e, doc)
            if bb is None:
                continue
            ext_w, ext_h = bb.bounds[2] - bb.bounds[0], bb.bounds[3] - bb.bounds[1]
            if not bb.intersects(vung):
                continue
            c = nd.theo_ten(e.dxf.name)
            if c is None:
                kids = con_insert(e.dxf.name)
                if len(kids) == 1 and len(list(doc.blocks.get(e.dxf.name))) == 1:
                    c = nd.theo_ten(kids[0].dxf.name)          # block bao ngoai (vd mirror) chua 1 block thiet bi
            shp = None
            if c is None and tpp.layer_goc(e.dxf.layer) in nd.lay_tb:
                shp = hinh_ky_hieu(e)
                if shp is not None:
                    w, h = shp.bounds[2] - shp.bounds[0], shp.bounds[3] - shp.bounds[1]
                    c = nd.theo_chu_ky(e.dxf.layer, w, h)
            if c is not None:
                shp = shp or hinh_ky_hieu(e)
                if shp is None or not vung.contains(shp.centroid):
                    continue
                out.append(dict(cat=c, ma=c["ma"], x=shp.centroid.x, y=shp.centroid.y, fp=shp, rot=e.dxf.rotation % 360,
                                block=e.dxf.name, layer=e.dxf.layer, handle=e.dxf.handle, insert=(e.dxf.insert.x, e.dxf.insert.y)))
                continue
            if tpp.layer_goc(e.dxf.layer) in nd.lay_tb:
                shp = shp or hinh_ky_hieu(e)
                if shp is not None and vung.contains(shp.centroid) and shp.area < 2.5e6:
                    la.append(dict(block=e.dxf.name, layer=e.dxf.layer, x=shp.centroid.x, y=shp.centroid.y,
                                   rong=shp.bounds[2] - shp.bounds[0], cao=shp.bounds[3] - shp.bounds[1]))
                continue
            if depth < 3 and (ext_w > 2500 or ext_h > 2500):    # chi di sau vao block lon (xref/cum)
                try:
                    duyet(e.virtual_entities(), depth + 1)
                except Exception:
                    pass

    duyet(msp, 0)
    return out, la


# ------------------------------------------------------------------------------------------------- nhan dien noi that
def doc_noi_that(msp, doc, cfg, vung):
    nt = cfg["noi_that"]
    out = []

    def con(name):
        b = doc.blocks.get(name)
        return [x for x in b if x.dxftype() == "INSERT"] if b else []

    def ten_loai(name):
        for k, v in nt.items():
            if khop(name, v.get("ten", [])):
                return k
        return None

    def duyet(ents, depth):
        for e in ents:
            if e.dxftype() != "INSERT":
                continue
            bb = bb_ins(e, doc)
            if bb is None:
                continue
            ext_w, ext_h = bb.bounds[2] - bb.bounds[0], bb.bounds[3] - bb.bounds[1]
            if not bb.intersects(vung):
                continue
            kids = con(e.dxf.name)
            n_moc = sum(1 for k in kids if khop(k.dxf.name, nt["tu_ao"]["con_moc_ao"]))
            n_tab = sum(1 for k in kids if khop(k.dxf.name, nt["giuong"]["con_tab_dau_giuong"]))
            loai = ten_loai(e.dxf.name)
            if n_moc >= nt["tu_ao"]["so_moc_ao_toi_thieu"]:
                loai = "tu_ao"
            elif n_tab >= 2 and loai is None:
                loai = "giuong"
            if loai and bb.area < 40e6:
                rec = dict(loai=loai, ten=ten_goc(e.dxf.name), fp=bb, x=bb.centroid.x, y=bb.centroid.y, handle=e.dxf.handle)
                if loai == "tu_ao":
                    moc = [m for m in e.virtual_entities() if m.dxftype() == "INSERT" and khop(m.dxf.name, nt["tu_ao"]["con_moc_ao"])]
                    rects = []
                    for m in moc:
                        try:
                            x2 = bbox.extents([m], fast=True)
                            rects.append(box(x2.extmin.x, x2.extmin.y, x2.extmax.x, x2.extmax.y))
                        except Exception:
                            pass
                    if rects:
                        g = unary_union(rects).buffer(150, join_style=2).intersection(bb)
                        rec["fp"] = g if not g.is_empty else bb
                elif loai == "giuong":
                    tabs = [m for m in e.virtual_entities() if m.dxftype() == "INSERT" and khop(m.dxf.name, nt["giuong"]["con_tab_dau_giuong"])]
                    tb = []
                    for m in tabs:
                        try:
                            x2 = bbox.extents([m], fast=True)
                            tb.append(box(x2.extmin.x, x2.extmin.y, x2.extmax.x, x2.extmax.y))
                        except Exception:
                            pass
                    rec.update(phan_tich_giuong(bb, tb, cfg["nguong"]["vung_goi_sau"]))
                out.append(rec)
                continue
            if depth < 3 and (ext_w > 2500 or ext_h > 2500):
                try:
                    duyet(e.virtual_entities(), depth + 1)
                except Exception:
                    pass

    duyet(msp, 0)
    return out


def phan_tich_giuong(bb, tabs, sau_goi):
    """Tu hop bao bo giuong va 2 tab dau giuong: than giuong, truc giuong, vung goi."""
    minx, miny, maxx, maxy = bb.bounds
    if len(tabs) < 2:
        return dict(truc=None, vung_goi=None, than=bb)
    a, b = tabs[0], tabs[1]
    if abs(a.centroid.y - b.centroid.y) < abs(a.centroid.x - b.centroid.x):     # tab nam ngang -> truc giuong doc (song song Y)
        lx, rx = sorted([a, b], key=lambda t: t.centroid.x)
        x0, x1 = lx.bounds[2], rx.bounds[0]
        dau_tren = (a.centroid.y + b.centroid.y) / 2 > bb.centroid.y
        than = box(x0, miny, x1, maxy)
        goi = box(x0, maxy - sau_goi, x1, maxy) if dau_tren else box(x0, miny, x1, miny + sau_goi)
        return dict(truc=("x", (x0 + x1) / 2), vung_goi=goi, than=than)
    lo, hi = sorted([a, b], key=lambda t: t.centroid.y)
    y0, y1 = lo.bounds[3], hi.bounds[1]
    dau_phai = (a.centroid.x + b.centroid.x) / 2 > bb.centroid.x
    than = box(minx, y0, maxx, y1)
    goi = box(maxx - sau_goi, y0, maxx, y1) if dau_phai else box(minx, y0, minx + sau_goi, y1)
    return dict(truc=("y", (y0 + y1) / 2), vung_goi=goi, than=than)


# --------------------------------------------------------------------------------------------------------- phong
def _lo_khep_gan_dung(chains, khung, r):
    """Khep khe <= 2r tren mang net (nong r roi lay lo ben trong) - ranh gan dung cho phong chua dong kin.
    Tinh mot lan: nong tung doan thang roi union (nhanh hon nhieu so voi union net roi moi nong); tra ve (ds lo, STRtree)."""
    segs = []
    for c, pts in chains:
        q = pts + ([pts[0]] if c else [])
        segs += [(q[i], q[i + 1]) for i in range(len(q) - 1) if q[i] != q[i + 1]]
    for k in khung:
        cs = list(k.coords)
        segs += [(cs[i], cs[i + 1]) for i in range(len(cs) - 1) if cs[i] != cs[i + 1]]
    if not segs:
        return [], shapely.STRtree([Polygon()])
    bufs = shapely.buffer(shapely.linestrings([[a, b] for a, b in segs]), r, quad_segs=1,
                          cap_style="square", join_style="mitre")
    g = shapely.union_all(bufs)
    lo = [Polygon(ring) for poly in getattr(g, "geoms", [g]) for ring in poly.interiors]
    lo = [h for h in lo if h.area > 0.5e6]
    return lo, shapely.STRtree(lo if lo else [Polygon()])


def dung_phong(msp, doc, cfg, args):
    chains = tpp.doc_chains(msp, cfg["layer_ranh"])
    khung, _ = tpp.doc_khung_cua(msp, chains, cfg["layer_cua"])
    faces, _, _ = tpp.build_faces(chains, 1200.0, khung)
    seeds = tpp.doc_seeds(msp, cfg["layer_ten_phong"], doc)
    re_can = re.compile(cfg["nhan_can_ho"], re.I)
    loai_re = [(k, re.compile(p, re.I)) for k, p in cfg["loai_phong"]]
    nhan_can, phong_seed = [], []
    for t, (x, y), info in seeds:
        tt = bo_dau(t)
        if re_can.match(tt):
            nhan_can.append((t.strip(), Point(x, y)))
            continue
        loai = [k for k, r in loai_re if r.search(tt)]
        if loai:
            phong_seed.append((t.strip(), Point(x, y), loai))
    # ranh gan dung cho phong chua dong kin (mat dung chi ve ngat quang...): khep khe tang dan, chi nhan lo chua DUNG
    # mot ten phong va khong qua lon
    ds_r = cfg["nguong"].get("dong_ranh_gan_dung_cac_buoc", [cfg["nguong"]["dong_ranh_gan_dung"]])
    dt_max = cfg["nguong"].get("dt_phong_gan_dung_toi_da", 60.0) * 1e6
    face_tree = shapely.STRtree(faces) if faces else None
    lo_dong = {}                         # r -> (cac lo kin sau khi khep khe <= 2r, STRtree); tinh khi can
    tat_ca_seed = [p for _, p, _ in phong_seed]
    rooms = {}
    for ten, p, loai in phong_seed:
        hit = [faces[i] for i in face_tree.query(p, predicate="within")] if face_tree is not None else []
        hit = [f for f in hit if f.area < 150e6]
        gan_dung = False
        if hit:
            f = min(hit, key=lambda g: g.area)
        else:
            f = None
            for r in ds_r:
                if r not in lo_dong:
                    lo_dong[r] = _lo_khep_gan_dung(chains, khung, r)
                for i in lo_dong[r][1].query(p, predicate="within"):
                    h = lo_dong[r][0][i]
                    if h.area > dt_max or sum(1 for q in tat_ca_seed if h.contains(q)) > 1:
                        continue
                    f = h.buffer(r, join_style=2, mitre_limit=3)
                    gan_dung = r
                    break
                if f is not None:
                    break
        if f is None or f.area > 150e6:
            rooms[("ho", ten, round(p.x), round(p.y))] = dict(ten=ten, loai=loai, poly=None, gan_dung=False, seed=p)
            continue
        key = (round(f.centroid.x), round(f.centroid.y), round(f.area))
        if key in rooms:
            rooms[key]["ten"] += " + " + ten
            rooms[key]["loai"] = sorted(set(rooms[key]["loai"] + loai))
        else:
            rooms[key] = dict(ten=ten, loai=loai, poly=f, gan_dung=gan_dung, seed=p)
    ds = list(rooms.values())
    cans = _nhom_can_ho(msp, ds, cfg, nhan_can)
    # danh so phong trung ten trong cung can (tren xuong duoi, trai sang phai) de bao cao phan biet duoc
    dem = defaultdict(list)
    for x in ds:
        dem[(x["can"], x["ten"])].append(x)
    for v in dem.values():
        if len(v) > 1:
            for k, x in enumerate(sorted(v, key=lambda x: (-round(x["seed"].y, -2), x["seed"].x)), 1):
                x["ten"] = f"{x['ten']} {k}"
    return ds, cans


def _nhom_can_ho(msp, ds, cfg, nhan_can):
    """Nhom phong thanh can ho qua CUA DI: hai phong cung cham mot cua (hop bao cua <= 4 m) thi cung can; tuong chung
    giua hai can khong co cua nen khong bi gop (nong ranh phong se dinh hai can qua tuong chung). Phong khong noi
    qua cua nao (vd vach kinh, cua lua) thi gop vao can ke ben co canh chung dai nhat. Ten can: ghep cap (cum, nhan
    CHxx) gan nhat truoc. Phong khong dung duoc ranh: gan theo cum chua diem ten phong."""
    idx = [i for i, x in enumerate(ds) if x["poly"] is not None]
    cha = {i: i for i in idx}

    def goc(i):
        while cha[i] != i:
            cha[i] = cha[cha[i]]
            i = cha[i]
        return i

    def noi(a, b):
        cha[goc(a)] = goc(b)

    want = {tpp.layer_goc(x) for x in cfg["layer_cua"]}
    tree = shapely.STRtree([ds[i]["poly"] for i in idx]) if idx else None
    for ents in tpp._nhom_cua(msp, want):
        doan, _ = tpp._doan_cua(ents)
        if not doan or tree is None:
            continue
        xs = [p[0] for s in doan for p in s]
        ys = [p[1] for s in doan for p in s]
        if max(max(xs) - min(xs), max(ys) - min(ys)) > 4000:
            continue
        hb = shapely.box(min(xs), min(ys), max(xs), max(ys)).buffer(50, join_style=2)
        cham = [idx[k] for k in tree.query(hb) if ds[idx[k]]["poly"].intersection(hb).area > 0.02e6]
        for a, b in zip(cham, cham[1:]):
            noi(a, b)
    # cum khong co phong khach (phong le, cum PN + WC khong nhan duoc cua...) -> gop vao cum ke ben co canh chung dai nhat
    doi = True
    while doi:
        doi = False
        for i in idx:
            cum_i = [j for j in idx if goc(j) == goc(i)]
            if any("khach" in ds[j]["loai"] for j in cum_i):
                continue
            vien = unary_union([ds[j]["poly"] for j in cum_i]).buffer(350, join_style=2)
            khac = [j for j in idx if goc(j) != goc(i)]
            tot = max(khac, key=lambda j: vien.intersection(ds[j]["poly"]).area, default=None)
            if tot is not None and vien.intersection(ds[tot]["poly"]).area > 0.1e6:
                noi(i, tot)
                doi = True
                break
    nhom = defaultdict(list)
    for i in idx:
        nhom[goc(i)].append(i)
    cums = [unary_union([ds[i]["poly"].buffer(150, join_style=2) for i in v]).buffer(-150, join_style=2) for v in nhom.values()]
    thanh_vien = list(nhom.values())
    # ghep ten: cap (cum, nhan) gan nhat truoc
    cap = sorted((c.distance(p), k, n) for k, c in enumerate(cums) for n, (_, p) in enumerate(nhan_can))
    ten_cum, da_dung = {}, set()
    for d, k, n in cap:
        if k in ten_cum or n in da_dung or d > 3000:
            continue
        ten_cum[k], _ = nhan_can[n][0], da_dung.add(n)
    thu_tu = sorted(range(len(cums)), key=lambda k: -cums[k].centroid.y)
    out, so = [], 0
    for k in thu_tu:
        if k not in ten_cum:
            so += 1
            ten_cum[k] = f"Căn {so}"
        for i in thanh_vien[k]:
            ds[i]["can"] = ten_cum[k]
        out.append((cums[k], ten_cum[k]))
    for x in ds:
        if x["poly"] is None:
            x["can"] = next((t for c, t in out if c.buffer(10).contains(x["seed"])), "?")
    return out


# ------------------------------------------------------------------------------------------------------ kiem tra
class SoTay:
    def __init__(self):
        self.ds = []

    def them(self, can, phong, tb, loai, hang_muc, mo_ta, gia_tri=None, nguong=None, de_xuat=None, trang_thai=None):
        self.ds.append(dict(can=can, phong=phong, tb=tb, loai=loai, hang_muc=hang_muc, mo_ta=mo_ta, gia_tri=gia_tri,
                            nguong=nguong, de_xuat=de_xuat, trang_thai=trang_thai or ("TECHNICAL REVIEW REQUIRED" if loai == CRIT else "CẦN SỬA")))


def kc_tuong(room_poly, x, y):
    return room_poly.exterior.distance(Point(x, y)) if room_poly.contains(Point(x, y)) else -room_poly.exterior.distance(Point(x, y))


def bo_tri_lai_den(room, dens, cam, ng, giuong):
    """Luoi den chung moi: cach tuong ~600 (toi thieu 500), tam-tam >= 1200, theo truc giuong neu co, khong vao vung cam.
    So den <= so den hien co (giam den chu khong ep khoang cach)."""
    N = len(dens)
    if N == 0:
        return []
    for off in (ng["den_cach_tuong_uu_tien"], ng["den_cach_tuong_toi_thieu"]):
        A = room["poly"].buffer(-off, join_style=2)
        if not A.is_empty:
            break
    if A.is_empty:
        return []
    A = A.difference(cam) if not cam.is_empty else A
    if A.is_empty:
        return []
    minx, miny, maxx, maxy = room["poly"].buffer(-off, join_style=2).bounds
    s = ng["den_kc_toi_thieu"]

    def truc_vi_tri(lo, hi, n, tam=None):
        if tam is not None:
            h = min(tam - lo, hi - tam)
            if n == 1:
                return [tam]
            if h <= 0 or 2 * h / (n - 1) < s:
                return None
            return [tam - h + i * 2 * h / (n - 1) for i in range(n)]
        if n == 1:
            return [(lo + hi) / 2]
        if (hi - lo) / (n - 1) < s:
            return None
        return [lo + i * (hi - lo) / (n - 1) for i in range(n)]

    tam_x = tam_y = None
    if giuong and giuong.get("truc"):
        if giuong["truc"][0] == "x":
            tam_x = giuong["truc"][1]
        else:
            tam_y = giuong["truc"][1]
    best = None
    for nx in range(1, 8):
        for ny in range(1, 8):
            if nx * ny > N:
                continue
            xs = truc_vi_tri(minx, maxx, nx, tam_x)
            ys = truc_vi_tri(miny, maxy, ny, tam_y)
            if xs is None or ys is None:
                continue
            pts = [(x, y) for x in xs for y in ys if A.buffer(1).contains(Point(x, y))]
            if giuong and giuong.get("vung_goi") is not None:
                pts = [q for q in pts if not giuong["vung_goi"].contains(Point(q))]
            sc = (len(pts), -abs(nx - ny))
            if pts and (best is None or sc > best[0]):
                best = (sc, pts)
    return best[1] if best else []


def soat(ds_phong, thiet_bi, noi_that, ng, so):
    de_xuat = []          # (thiet bi cu hoac None, ma, x, y, rot, ly do)
    phong_kq = []
    for r in ds_phong:
        if r["poly"] is None:
            so.them(r["can"], r["ten"], "", COORD, "Ranh phòng", "Không dựng được ranh phòng (nét tường/vách hở lớn); thiết bị trong phòng chưa soát được.")
            phong_kq.append(dict(r, tb=[], trang_thai="TECHNICAL REVIEW REQUIRED", truc="?"))
            continue
        P = r["poly"]
        tb = [d for d in thiet_bi if P.contains(Point(d["x"], d["y"]))]
        for d in tb:
            d["can"], d["phong"] = r["can"], r["ten"]
        nt = [f for f in noi_that if P.buffer(200).contains(Point(f["x"], f["y"])) or P.intersection(f["fp"]).area > 0.5 * f["fp"].area]
        tu = [f for f in nt if f["loai"] == "tu_ao"]
        giuong = next((f for f in nt if f["loai"] == "giuong"), None)
        ban = [f for f in nt if f["loai"] in ("ban_an",)]
        sofa = [f for f in nt if f["loai"] == "sofa"]
        n0 = len(so.ds)
        ten = r["ten"]
        truc = "trục hình học phòng"
        if "ngu" in r["loai"] and giuong and giuong.get("truc"):
            truc = f"trục giường ({giuong['truc'][0].upper()} = {giuong['truc'][1]:.0f})"
        elif "wc" in r["loai"]:
            truc = "trục thiết bị vệ sinh"
        elif set(r["loai"]) & {"an", "khach"} and ban:
            truc = "tâm bàn ăn / bộ sofa"
        if r["gan_dung"]:
            so.them(r["can"], ten, "", DESIGN, "Ranh phòng gần đúng",
                    f"Nền chưa khép kín phòng (mặt dựng/cửa sổ/vách kính vẽ ngắt quãng): ranh dựng gần đúng bằng cách khép khe hở "
                    f"≤ {2 * r['gan_dung'] / 1000:.1f} m; khoảng cách tới tường ở phía khe hở chỉ là ước lượng, cần kiểm tra lại.")
        # 1) tu ao
        cam_tu = unary_union([f["fp"] for f in tu]) if tu else Polygon()
        for d in tb:
            if not cam_tu.is_empty and d["fp"].intersects(cam_tu):
                if d["cat"]["he_thong"] in PCCC:
                    so.them(r["can"], ten, d, CRIT, "Thiết bị PCCC trong vùng tủ áo",
                            f"{d['ma']} nằm trong vùng tủ áo. Không tự dời/xóa thiết bị PCCC: cần tư vấn PCCC xem lại.")
                    d["st"] = "TECHNICAL REVIEW REQUIRED"
                else:
                    d["st"] = "CONFLICT"
                    d["can_doi"] = "tủ áo"
        # 2) den chung: khoang cach, cach tuong
        dens = [d for d in tb if d["cat"]["nhom"] == "den_chung"]
        vp_den = False
        for i, a in enumerate(dens):
            for b in dens[i + 1:]:
                kc = math.dist((a["x"], a["y"]), (b["x"], b["y"]))
                if kc < ng["den_kc_toi_thieu"] - 1:
                    so.them(r["can"], ten, a, HARD, "Khoảng cách đèn",
                            f"{a['ma']} cách {b['ma']} tại ({b['x']:.0f}, {b['y']:.0f}) {kc:.0f} mm < {ng['den_kc_toi_thieu']} mm.",
                            round(kc), ng["den_kc_toi_thieu"])
                    vp_den = True
        for a in dens:
            kt = kc_tuong(P, a["x"], a["y"])
            if kt < ng["den_cach_tuong_toi_thieu"] - 1:
                so.them(r["can"], ten, a, HARD, "Đèn cách tường", f"{a['ma']} cách tường {kt:.0f} mm < {ng['den_cach_tuong_toi_thieu']} mm.",
                        round(kt), ng["den_cach_tuong_toi_thieu"])
                vp_den = True
        # 3) chong lan giua thiet bi
        for i, a in enumerate(tb):
            for b in tb[i + 1:]:
                if a["fp"].intersects(b["fp"]) and a["fp"].intersection(b["fp"]).area > 1:
                    if a["ma"] == b["ma"] and math.dist((a["x"], a["y"]), (b["x"], b["y"])) < 50:
                        so.them(r["can"], ten, b, COORD, "Thiết bị chèn trùng",
                                f"Hai block {a['ma']} trùng vị trí (handle {a['handle']} và {b['handle']}): xóa bản trùng "
                                f"(OVERKILL) sau khi bộ môn xác nhận, kiểm tra lại số lượng trong bảng thống kê.")
                        continue
                    lo = a if a["cat"]["uu_tien"] > b["cat"]["uu_tien"] else b
                    hi = b if lo is a else a
                    giu = lo["cat"]["he_thong"] in PCCC or lo["cat"]["nhom"] == "den_tha"
                    so.them(r["can"], ten, lo, COORD, "Thiết bị chồng nhau",
                            f"{a['ma']} chồng lên {b['ma']}; " + (
                                f"{lo['ma']} giữ vị trí (PCCC / đèn thả theo tâm bàn ăn - sofa): cần phối hợp dời {hi['ma']} hoặc chấp nhận."
                                if giu else f"dời thiết bị ưu tiên thấp hơn ({lo['ma']}, {lo['cat']['uu_tien']})."))
                    if not giu:
                        lo["can_doi"] = lo.get("can_doi") or f"chồng {hi['ma']}"
        # 4) phong ngu
        if "ngu" in r["loai"] and giuong:
            for d in tb:
                if giuong.get("vung_goi") is not None and d["cat"]["nhom"] == "den_chung" and giuong["vung_goi"].contains(Point(d["x"], d["y"])):
                    so.them(r["can"], ten, d, DESIGN, "Đèn trên vùng gối", f"{d['ma']} nằm trên vùng gối (đầu giường {ng['vung_goi_sau']} mm).")
                    d["tren_goi"] = True
                if d["cat"]["nhom"] == "gio_cap" and giuong.get("vung_goi") is not None and d["fp"].distance(giuong["vung_goi"]) < 600:
                    so.them(r["can"], ten, d, COORD, "Gió cấp gần vùng gối",
                            f"Miệng gió cấp cách vùng gối {d['fp'].distance(giuong['vung_goi']):.0f} mm: có thể thổi thẳng vào đầu giường, cần HVAC xem hướng thổi.")
                if d["cat"]["nhom"] == "lo_tham" and d["fp"].intersects(giuong.get("than", giuong["fp"])):
                    so.them(r["can"], ten, d, DESIGN, "Lỗ thăm trên giường", "Lỗ thăm trần nằm trên giường; nên dời về phía cửa vào / vùng phụ.")
                    d["can_doi"] = d.get("can_doi") or "trên giường"
        # 5) WC
        if "wc" in r["loai"]:
            fx = [f for f in nt if f["loai"] in ("bon_cau", "chau_rua", "guong", "sen_tam", "vach_tam")]
            guong = [f for f in nt if f["loai"] == "guong"]
            truc_x = {round(f["x"]) for f in fx}
            truc_y = {round(f["y"]) for f in fx}
            for d in tb:
                if d["cat"]["nhom"] == "den_guong" and guong:
                    g = min(guong, key=lambda f: math.dist((f["x"], f["y"]), (d["x"], d["y"])))
                    gx0, gy0, gx1, gy1 = g["fp"].bounds
                    doc_x = (gx1 - gx0) >= (gy1 - gy0)           # guong nam ngang -> truc guong la x = tam
                    lech = abs(d["x"] - g["x"]) if doc_x else abs(d["y"] - g["y"])
                    if lech > ng["lech_truc_cho_phep"]:
                        nx, ny = (g["x"], d["y"]) if doc_x else (d["x"], g["y"])
                        so.them(r["can"], ten, d, DESIGN, "Đèn rọi gương lệch trục gương", f"Lệch trục gương/chậu rửa {lech:.0f} mm.",
                                round(lech), ng["lech_truc_cho_phep"], (nx, ny))
                        de_xuat.append((d, d["ma"], nx, ny, d["rot"], "theo trục gương"))
                if d["cat"]["nhom"] == "den_chung" and fx:
                    ok = any(abs(d["x"] - t) <= ng["lech_truc_cho_phep"] for t in truc_x) or any(abs(d["y"] - t) <= ng["lech_truc_cho_phep"] for t in truc_y)
                    if not ok:
                        so.them(r["can"], ten, d, DESIGN, "Đèn WC không theo trục thiết bị",
                                "Không trùng trục chậu rửa / bồn cầu / sen (sai lệch > 50 mm theo cả X và Y).")
            if not fx:
                so.them(r["can"], ten, "", DESIGN, "Thiếu dữ liệu WC", "Không nhận diện được thiết bị vệ sinh để kiểm tra trục.")
        # 6) den tha
        for d in tb:
            if d["cat"]["nhom"] != "den_tha":
                continue
            ref = ban + sofa
            if not ref:
                so.them(r["can"], ten, d, DESIGN, "Đèn thả chưa đối chiếu", "Không thấy bàn ăn/sofa trong phòng để đối chiếu vị trí đèn thả.")
                continue
            f = min(ref, key=lambda t: math.dist((t["x"], t["y"]), (d["x"], d["y"])))
            lech = math.dist((f["x"], f["y"]), (d["x"], d["y"]))
            if lech > ng["den_tha_lech_toi_da"]:
                so.them(r["can"], ten, d, DESIGN, "Đèn thả lệch tâm", f"Lệch tâm {('bàn ăn' if f['loai'] == 'ban_an' else 'bộ sofa')} {lech:.0f} mm.",
                        round(lech), ng["den_tha_lech_toi_da"], (f["x"], f["y"]))
                de_xuat.append((d, d["ma"], f["x"], f["y"], d["rot"], "tâm " + ("bàn ăn" if f["loai"] == "ban_an" else "bộ sofa")))
        # 7) lo tham giua phong khach
        if "khach" in r["loai"]:
            minx, miny, maxx, maxy = P.bounds
            giua = box(minx + (maxx - minx) / 3, miny + (maxy - miny) / 3, maxx - (maxx - minx) / 3, maxy - (maxy - miny) / 3)
            for d in tb:
                if d["cat"]["nhom"] == "lo_tham" and giua.contains(Point(d["x"], d["y"])):
                    so.them(r["can"], ten, d, DESIGN, "Lỗ thăm giữa phòng khách", "Nên dời về vùng biên / trần phụ nếu vẫn tiếp cận được thiết bị.")
        # 8) dau bao gan gio cap
        caps = [d for d in tb if d["cat"]["nhom"] == "gio_cap"]
        for d in tb:
            if d["cat"]["he_thong"] == "FIRE_ALARM" and caps:
                kc = min(c["fp"].distance(Point(d["x"], d["y"])) for c in caps)
                if kc < ng["dau_bao_cach_gio_cap"]:
                    so.them(r["can"], ten, d, COORD, "Đầu báo gần miệng gió cấp",
                            f"Cách miệng gió cấp {kc:.0f} mm (ngưỡng tạm {ng['dau_bao_cach_gio_cap']} mm, chưa chốt): cần tư vấn PCCC/HVAC xác nhận.",
                            round(kc), ng["dau_bao_cach_gio_cap"])
        # ---- de xuat ---------------------------------------------------------------------------------------
        co_dinh = [d for d in tb if d["cat"]["nhom"] != "den_chung"]
        cam = unary_union([cam_tu.buffer(100)] + [d["fp"].buffer(ng["vung_tranh_quanh_thiet_bi"]) for d in co_dinh]) if (co_dinh or not cam_tu.is_empty) else Polygon()
        if vp_den or any(d.get("st") == "CONFLICT" or d.get("can_doi") or d.get("tren_goi") for d in dens):
            moi = bo_tri_lai_den(r, dens, cam, ng, giuong if "ngu" in r["loai"] else None)
            if moi:
                for d in dens:
                    d["thay_the"] = True
                for (x, y) in moi:
                    de_xuat.append((None, dens[0]["ma"], x, y, 0.0, f"bố trí lại đèn chung {r['ten']}"))
                so.them(r["can"], ten, "", HARD if vp_den else DESIGN, "Đề xuất bố trí lại đèn",
                        f"{len(dens)} đèn hiện có → {len(moi)} đèn đề xuất (cách tường ≥ {ng['den_cach_tuong_toi_thieu']}, tâm-tâm ≥ {ng['den_kc_toi_thieu']}, theo {truc}).",
                        len(dens), None, None, "ĐỀ XUẤT")
            else:
                so.them(r["can"], ten, "", HARD, "Không đề xuất được lưới đèn",
                        "Phòng quá hẹp hoặc vướng thiết bị: không bố trí được lưới đèn thỏa 1200/500; cần giảm số đèn hoặc đổi giải pháp chiếu sáng.")
        for d in tb:
            if d["cat"]["nhom"] == "den_chung" or d["cat"]["he_thong"] in PCCC or not d.get("can_doi"):
                continue
            half = max(d["fp"].bounds[2] - d["fp"].bounds[0], d["fp"].bounds[3] - d["fp"].bounds[1]) / 2
            khac = [x["fp"].buffer(half + ng["vung_tranh_quanh_thiet_bi"] / 2) for x in tb if x is not d]
            A = P.buffer(-(half + 100), join_style=2).difference(unary_union([cam_tu.buffer(100)] + khac))
            if giuong and d["cat"]["nhom"] == "lo_tham":
                A = A.difference(giuong.get("than", giuong["fp"]).buffer(100))
            q = nearest_points(A, Point(d["x"], d["y"]))[0] if not A.is_empty else None
            if q is None or math.dist((q.x, q.y), (d["x"], d["y"])) < 10:
                so.them(r["can"], ten, d, COORD, "Không tìm được vị trí thay thế",
                        f"{d['ma']} ({d['can_doi']}): không còn chỗ trống thỏa điều kiện trong phòng, cần thiết kế xem lại.")
                continue
            de_xuat.append((d, d["ma"], q.x, q.y, d["rot"], f"ra khỏi vùng {d['can_doi']}"))
            so.them(r["can"], ten, d, COORD if d.get("st") == "CONFLICT" else DESIGN, "Dời thiết bị",
                    f"{d['ma']} ({d['can_doi']}) → đề xuất ({q.x:.0f}, {q.y:.0f}).", None, None, (q.x, q.y))
        # trang thai phong
        moi_loai = {x["loai"] for x in so.ds[n0:]}
        st = "TECHNICAL REVIEW REQUIRED" if CRIT in moi_loai else "REVISE" if HARD in moi_loai else \
            "PASS WITH WARNING" if moi_loai else "PASS"
        phong_kq.append(dict(r, tb=tb, trang_thai=st, truc=truc, n_tu=len(tu), giuong=bool(giuong)))
    return phong_kq, de_xuat


# ---------------------------------------------------------------------------------------------------- dau ra
def xuat_excel(path, du_an, file, phong_kq, so, thiet_bi, de_xuat, chua_biet):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Tom tat"
    ws.append(["Căn hộ", "Phòng", "Loại", "Trục thiết kế chính", "Số thiết bị", "CRITICAL", "HARD-RULE", "COORDINATION", "DESIGN", "Trạng thái"])
    for p in phong_kq:
        c = defaultdict(int)
        for x in so.ds:
            if x["can"] == p["can"] and x["phong"] == p["ten"]:
                c[x["loai"]] += 1
        ws.append([p["can"], p["ten"], ", ".join(p["loai"]), p.get("truc", ""), len(p.get("tb", [])), c[CRIT], c[HARD], c[COORD], c[DESIGN], p["trang_thai"]])
    ws2 = wb.create_sheet("Loi va de xuat")
    ws2.append(["STT", "Dự án", "File", "Căn hộ", "Phòng", "Thiết bị", "Handle", "X", "Y", "Hạng mục", "Mô tả", "Giá trị đo (mm)",
                "Ngưỡng (mm)", "Loại cảnh báo", "Mức độ", "Đề xuất X", "Đề xuất Y", "Người phụ trách", "Trạng thái"])
    for i, x in enumerate(so.ds, 1):
        d = x["tb"] if isinstance(x["tb"], dict) else None
        ws2.append([i, du_an, file, x["can"], x["phong"], d["ma"] if d else "", d["handle"] if d else "",
                    round(d["x"]) if d else None, round(d["y"]) if d else None, x["hang_muc"], x["mo_ta"], x["gia_tri"], x["nguong"],
                    x["loai"], MUC[x["loai"]], round(x["de_xuat"][0]) if x["de_xuat"] else None,
                    round(x["de_xuat"][1]) if x["de_xuat"] else None, "", x["trang_thai"]])
    ws3 = wb.create_sheet("Thiet bi")
    ws3.append(["DEVICE_ID", "ROOM", "APARTMENT", "DEVICE_TYPE", "BLOCK_NAME (thư viện)", "BLOCK gốc", "X", "Y", "Z", "ROTATION", "WIDTH", "HEIGHT",
                "SYSTEM", "PRIORITY", "STATUS", "WARNING"])
    dem = defaultdict(int)
    for d in sorted(thiet_bi, key=lambda t: (t.get("can", "~"), t.get("phong", "~"), t["ma"])):
        dem[(d.get("phong"), d["ma"])] += 1
        warns = [x["hang_muc"] for x in so.ds if x["tb"] is d]
        st = d.get("st") or ("WARN" if warns else "PASS")
        ws3.append([f"{bo_dau(d.get('phong') or 'ngoai').upper().replace(' ', '')[:10]}_{d['ma']}_{dem[(d.get('phong'), d['ma'])]:02d}",
                    d.get("phong", "(ngoài phòng)"), d.get("can", ""), d["cat"]["ten"], d["ma"], ten_goc(d["block"]), round(d["x"]), round(d["y"]),
                    "CEILING_LEVEL", round(d["rot"]), round(d["fp"].bounds[2] - d["fp"].bounds[0]), round(d["fp"].bounds[3] - d["fp"].bounds[1]),
                    d["cat"]["he_thong"], d["cat"]["uu_tien"], st, "; ".join(warns) or "NONE"])
    ws4 = wb.create_sheet("De xuat vi tri")
    ws4.append(["STT", "Thiết bị", "BLOCK_NAME", "Từ X", "Từ Y", "Đến X", "Đến Y", "Rotation", "Lý do"])
    for i, (d, ma, x, y, rot, ly) in enumerate(de_xuat, 1):
        ws4.append([i, ma, ma, round(d["x"]) if d else None, round(d["y"]) if d else None, round(x), round(y), round(rot), ly])
    if chua_biet:
        ws5 = wb.create_sheet("Block chua nhan dien")
        ws5.append(["Block", "Layer", "X", "Y", "Rộng", "Cao"])
        for c in chua_biet:
            ws5.append([c["block"], c["layer"], round(c["x"]), round(c["y"]), round(c["rong"]), round(c["cao"])])
    mau = {CRIT: "F8CBAD", HARD: "FCE4D6", COORD: "FFF2CC", DESIGN: "E2EFDA"}
    for row in ws2.iter_rows(min_row=2):
        row[13].fill = PatternFill("solid", fgColor=mau.get(row[13].value, "FFFFFF"))
    for w in wb.worksheets:
        for c in w[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for col in w.columns:
            L = max(len(str(c.value or "")) for c in col)
            w.column_dimensions[col[0].column_letter].width = min(60, max(8, L * 0.9))
    wb.save(path)


def ve_anh(path, can_poly, phong_kq, thiet_bi, noi_that, de_xuat, so):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    minx, miny, maxx, maxy = can_poly.bounds
    fig = plt.figure(figsize=(12, 12 * (maxy - miny + 1) / (maxx - minx + 1)), dpi=100)
    ax = fig.add_axes([0.02, 0.02, 0.96, 0.96])
    for p in phong_kq:
        if p["poly"] is None or not can_poly.buffer(10).contains(p["poly"].representative_point()):
            continue
        xs, ys = p["poly"].exterior.xy
        col = {"PASS": "#d3f9d8", "PASS WITH WARNING": "#fff3bf", "REVISE": "#ffc9c9", "TECHNICAL REVIEW REQUIRED": "#ffa8a8"}[p["trang_thai"]]
        ax.fill(xs, ys, color=col, alpha=0.6, lw=0)
        ax.plot(xs, ys, color="#495057", lw=1, ls="--" if p["gan_dung"] else "-")
        c = p["poly"].representative_point()
        ax.text(c.x, c.y, f"{p['ten']}\n{p['trang_thai']}", fontsize=7, ha="center", color="#212529")
    for f in noi_that:
        if not can_poly.buffer(300).contains(Point(f["x"], f["y"])):
            continue
        g = f["fp"]
        for gg in getattr(g, "geoms", [g]):
            xs, ys = gg.exterior.xy
            if f["loai"] == "tu_ao":
                ax.fill(xs, ys, color="#868e96", alpha=0.45, hatch="////", lw=0)
            elif f["loai"] == "giuong":
                ax.plot(xs, ys, color="#1971c2", lw=0.8)
        if f["loai"] == "giuong" and f.get("vung_goi") is not None:
            xs, ys = f["vung_goi"].exterior.xy
            ax.fill(xs, ys, color="#74c0fc", alpha=0.35, lw=0)
    loi_tb = {id(x["tb"]) for x in so.ds if isinstance(x["tb"], dict) and x["loai"] in (CRIT, HARD, COORD)}
    for d in thiet_bi:
        if not can_poly.buffer(10).contains(Point(d["x"], d["y"])):
            continue
        xs, ys = d["fp"].exterior.xy
        col = "#e03131" if id(d) in loi_tb else "#2f9e44"
        ax.plot(xs, ys, color=col, lw=1.2)
        if d.get("thay_the"):
            ax.plot(d["x"], d["y"], "x", color="#e03131", ms=6)
    for d, ma, x, y, rot, ly in de_xuat:
        if not can_poly.buffer(10).contains(Point(x, y)):
            continue
        ax.plot(x, y, "o", mfc="none", mec="#1c7ed6", ms=9, mew=2)
        if d is not None:
            ax.annotate("", xy=(x, y), xytext=(d["x"], d["y"]), arrowprops=dict(arrowstyle="->", color="#1c7ed6"))
    ax.set_aspect("equal")
    ax.set_xlim(minx - 300, maxx + 300)
    ax.set_ylim(miny - 300, maxy + 300)
    ax.set_title("Đỏ: thiết bị lỗi | Xanh lá: đạt | Vòng xanh dương: vị trí đề xuất | X đỏ: đèn thay bằng lưới mới | Xám gạch: tủ áo | Xanh nhạt: vùng gối", fontsize=8)
    fig.savefig(path)
    plt.close(fig)


def xuat_scr(path, so, de_xuat, thu_vien, cfg, catalog):
    lay_loi = cfg["layer_loi"]
    hau_to = cfg["hau_to_layer_de_xuat"]
    lay_cat = {c["ma"]: c["layer"] for c in catalog["thiet_bi"]}
    def lisp_str(s):
        return '"' + tpp.acad_str(s).replace("\\", "\\\\").replace('"', "'") + '"'

    # layer tao/dat bang LISP: ten layer co dau cach (A-Hoan thien tran-DX) se bi script hieu la Enter neu go qua -LAYER
    L = ["CMDECHO", "0", "OSMODE", "0", "ATTREQ", "0", "FILEDIA", "0",
         "(defun mk-lay (n c) (if (not (tblsearch \"LAYER\" n)) (entmake (list (cons 0 \"LAYER\") (cons 100 \"AcDbSymbolTableRecord\") "
         "(cons 100 \"AcDbLayerTableRecord\") (cons 2 n) (cons 70 0) (cons 62 c) (cons 6 \"Continuous\")))) n)",
         f"(mk-lay {lisp_str(lay_loi)} 1)"]
    for lay in sorted(set(lay_cat.values())):
        L += [f"(mk-lay {lisp_str(lay + hau_to)} 3)"]
    n = 0
    for x in so.ds:
        d = x["tb"] if isinstance(x["tb"], dict) else None
        if d is None or x["loai"] == DESIGN and not x["de_xuat"]:
            continue
        n += 1
        # entmake (khong qua lenh): khong phu thuoc OSNAP / text style co chieu cao co dinh; cat chuoi TRUOC khi ma hoa \U+
        nhan = lisp_str(f"L{n:02d} {x['hang_muc'][:40]}")
        L += [f'(entmake (list (cons 0 "CIRCLE") (cons 8 {lisp_str(lay_loi)}) (list 10 {d["x"]:.1f} {d["y"]:.1f} 0.0) (cons 40 250.0)))',
              f'(entmake (list (cons 0 "TEXT") (cons 8 {lisp_str(lay_loi)}) (list 10 {d["x"] + 280:.1f} {d["y"] - 60:.1f} 0.0) '
              f'(cons 40 120.0) (cons 1 {nhan})))']
    da_nap = set()
    for d, ma, x, y, rot, ly in de_xuat:
        f = os.path.join(thu_vien, ma + ".dwg").replace("\\", "/")
        if " " in f:
            raise SystemExit(f"Đường dẫn thư viện có dấu cách, -INSERT trong script sẽ hỏng: {f}")
        L += [f"(setvar \"CLAYER\" {lisp_str(lay_cat.get(ma, '0') + hau_to)})"]
        name = ma if ma in da_nap else f"{ma}={f}"
        da_nap.add(ma)
        L += ["_.-INSERT", name, "_S", "1", "_R", f"{rot:.2f}", f"{x:.1f},{y:.1f}"]
        if d is not None:
            L += [f'(entmake (list (cons 0 "LINE") (cons 8 {lisp_str(lay_loi)}) (list 10 {d["x"]:.1f} {d["y"]:.1f} 0.0) (list 11 {x:.1f} {y:.1f} 0.0)))']
    L += ['(setvar "CLAYER" "0")', "_.QSAVE", "_.QUIT", "_Y"]
    with open(path, "w", encoding="ascii", errors="replace", newline="\n") as fh:
        fh.write("\n".join(L) + "\n")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("dxf")
    ap.add_argument("--out-dir", default=".")
    ap.add_argument("--du-an", default="")
    ap.add_argument("--cau-hinh", default=os.path.join(HERE, "cau_hinh_tran.json"))
    ap.add_argument("--catalog", default=os.path.join(SKILL, "assets", "thu-vien", "catalog.json"))
    ap.add_argument("--thu-vien", default=os.path.join(SKILL, "assets", "thu-vien"))
    a = ap.parse_args()
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    cfg = json.load(open(a.cau_hinh, encoding="utf-8"))
    catalog = json.load(open(a.catalog, encoding="utf-8"))
    ng = cfg["nguong"]
    tpp._BO_QUA[0] = tpp.BLOCK_BO_QUA
    doc = ezdxf.readfile(a.dxf)
    msp = doc.modelspace()
    van_de = []
    if doc.header.get("$INSUNITS", 0) != 4:
        van_de.append(f"INSUNITS={doc.header.get('$INSUNITS', 0)}, chuẩn là 4 (mm).")
    t0 = time.time()
    ds_phong, cans = dung_phong(msp, doc, cfg, a)
    print(f'[phong] {time.time() - t0:.0f}s', file=sys.stderr)
    if not cans or all(c.is_empty for c, _ in cans):
        raise SystemExit("Không dựng được phòng nào: kiểm tra layer tên phòng / nét tường (cau_hinh_tran.json).")
    vung = unary_union([c for c, _ in cans]).buffer(300)
    nd = NhanDien(catalog, cfg)
    thiet_bi, chua_biet = doc_thiet_bi(msp, doc, nd, vung)
    print(f'[thiet bi] {time.time() - t0:.0f}s', file=sys.stderr)
    noi_that = doc_noi_that(msp, doc, cfg, vung)
    print(f'[noi that] {time.time() - t0:.0f}s', file=sys.stderr)
    so = SoTay()
    phong_kq, de_xuat = soat(ds_phong, thiet_bi, noi_that, ng, so)
    ngoai = [d for d in thiet_bi if not d.get("phong")]
    for d in ngoai:
        so.them("?", "(ngoài phòng đã dựng)", d, DESIGN, "Thiết bị ngoài phòng", f"{d['ma']} không nằm trong phòng nào đã dựng ranh (hành lang chung, phòng hở…).")
    os.makedirs(a.out_dir, exist_ok=True)
    xlsx = os.path.join(a.out_dir, "BaoCaoSoatTran.xlsx")
    xuat_excel(xlsx, a.du_an, os.path.basename(a.dxf), phong_kq, so, thiet_bi, de_xuat, chua_biet)
    anh = []
    for poly, ten in cans:
        p = os.path.join(a.out_dir, f"xem_tran_{bo_dau(ten).replace(' ', '_')}.png")
        ve_anh(p, poly, phong_kq, thiet_bi, noi_that, de_xuat, so)
        anh.append(p)
    scr = os.path.join(a.out_dir, "ve_de_xuat_tran.scr")
    xuat_scr(scr, so, de_xuat, a.thu_vien, cfg, catalog)
    dem = defaultdict(int)
    for x in so.ds:
        dem[x["loai"]] += 1
    kq = dict(
        file=os.path.basename(a.dxf), van_de=van_de,
        can_ho=[dict(ten=t, so_phong=sum(1 for p in phong_kq if p["can"] == t)) for _, t in cans],
        phong=[dict(can=p["can"], ten=p["ten"], loai=p["loai"], trang_thai=p["trang_thai"], truc=p.get("truc"),
                    ranh_gan_dung=p["gan_dung"], so_thiet_bi=len(p.get("tb", [])),
                    dien_tich=round(p["poly"].area / 1e6, 2) if p["poly"] is not None else None) for p in phong_kq],
        thiet_bi_theo_ma={m: sum(1 for d in thiet_bi if d["ma"] == m) for m in sorted({d["ma"] for d in thiet_bi})},
        noi_that_theo_loai={m: sum(1 for f in noi_that if f["loai"] == m) for m in sorted({f["loai"] for f in noi_that})},
        block_chua_nhan_dien=len(chua_biet), canh_bao=dict(dem), so_de_xuat=len(de_xuat),
        xlsx=xlsx, anh=anh, scr=scr)
    print(json.dumps(kq, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
