#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Mat bang cap dien O CAM can ho Archivina (skill cap-dien-ch): BO TRI MOI va SOAT ban ve co san.

Nguyen tac (nguoi dung chot 07/10/2026, references/nguyen-tac-bo-tri.md):
  - Thiet bi bo tri THEO NOI THAT (thieu noi that / ten phong -> DUNG LAI HOI, ghi vao 'can_hoi').
  - P. ngu: 2 o G hai ben dau giuong cach mep giuong 200-250; o TV thang tam tivi bam tuong; 1 o doi thuong gan cua
    phong; o ban lam viec. P. khach: TV + DN canh nhau (100 de vuong / 150 de chu nhat) theo tam tivi; 2 o hai dau sofa.
    Bep: o B tren mat bep (SO LUONG HOI NGUOI DUNG), BT duoi bep tu, HM theo tam may hut mui, TL sau tu lanh; may rua
    bat / lo nuong hoi truoc. WC: W canh lavabo ~300, X canh bon cau 150, BNL trong WC, cong tac 20A + cong tac 3 phim
    ngoai cua phia tay nam. Lo gia: W chong am cho may giat. Hop AC tai dan nong (cach tran 300). TD-CH cach khuon cua
    500 (tam tu). VDP: ngoai cua chuong/camera, trong can man hinh.
  - Cach khuon cua >= 200 (tam); khong sau canh cua, sau tu ao; uu tien tuong xay, tranh vach BTCT (khong tranh duoc ->
    danh dau may 'dat cho khi do cot vach').
  - Kich thuoc: tu mep tuong den tam thiet bi; thiet bi gan nhau -> chuoi kich thuoc giua 2 thiet bi (khong chong net).
  - Chia lo: khach + WC chung 1 lo; bep 1 lo (gom tu lanh, may giat); cac PN 1 lo (> 3 PN tach 2 lo); AC, BT, moi BNL 1
    lo rieng. F1, F2.. = quat hut WC.

    python cap_dien.py bo-tri <file.dxf> --out-dir <thu muc> [--so-o-bep 2 | "CH01=2,CH02=3"] [--de-o chu_nhat|vuong]
                       [--can CH01,CH02] [--may-rua-bat] [--lo-nuong] [--tv-pn-theo-truc-giuong] [--nhan-dien-bo-sung f.json]
    python cap_dien.py soat   <file.dxf> --out-dir <thu muc> [cac tuy chon nhu tren]

Xuat: JSON tom tat (stdout) + bo_tri_o_cam.json / soat_o_cam.json, BaoCao*.xlsx, xem_o_cam_<can>.png, *.scr (ve vao BAN SAO).
Khong sua DXF dau vao.
"""
import argparse
import fnmatch
import io
import json
import math
import os
import re
import sys
import time
from collections import Counter, defaultdict

import ezdxf
import shapely
from shapely.geometry import LineString, Point, Polygon, box
from shapely.geometry.polygon import orient
from shapely.ops import unary_union
from shapely.prepared import prep

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
_tran = os.environ.get("THIET_KE_TRAN_CH_SCRIPTS") or os.path.join(os.path.dirname(SKILL), "thiet-ke-tran-ch", "scripts")
sys.path.insert(0, _tran)
import soat_tran as st  # noqa: E402  (bo dung phong / noi that dung chung, kem tao_polyline_phong cua dien-tich-ch)
from nen_revit import dem_canh_cua  # noqa: E402

tpp = st.tpp
LOI, CB, GY = "Lỗi", "Cảnh báo", "Gợi ý"
THU_VIEN = os.path.join(SKILL, "assets", "thu-vien")


def bo_dau(s):
    return st.bo_dau(s)


def deg(v):
    return math.degrees(math.atan2(v[1], v[0]))


# ================================================================================================= ghi chu / cau hoi
class SoTay:
    """Ghi chu theo chuan bao cao Archivina: Lỗi / Cảnh báo / Gợi ý; 'hoi' = can hoi nguoi dung truoc khi ve."""

    def __init__(self):
        self.ds = []

    def them(self, can, phong, muc, hang_muc, mo_ta, hoi=False, tb=None, de_xuat=None, gia_tri=None, nguong=None):
        self.ds.append(dict(can=can or "", phong=phong or "", muc=muc, hang_muc=hang_muc, mo_ta=mo_ta, hoi=hoi, tb=tb,
                            de_xuat=de_xuat, gia_tri=gia_tri, nguong=nguong))


# ================================================================================================= doc ban ve
def doc_cua(msp, cfg, rooms):
    """Cua di co canh mo (ARC ban kinh 500-1200 tren layer cua, ke ca trong block/xref, ke ca net Revit da no).
    Moi cua: ban le H, ban kinh R, dau 'dong' (doc tuong) E, dau 'mo', phong chua cung quay."""
    want = {tpp.layer_goc(x) for x in cfg["layer_cua"]}
    want_cung = want | {tpp.layer_goc(x) for x in cfg.get("layer_cung_cua_them", [])}   # cung mo cua ve o layer khac (vd Revit: nét khuất)
    edges = []
    for r in rooms:
        if r["poly"] is None:
            continue
        c = list(r["poly"].exterior.coords)
        edges += [(a, b) for a, b in zip(c, c[1:]) if math.dist(a, b) > 200]
    tree = shapely.STRtree([LineString(e) for e in edges]) if edges else None
    polys = [r["poly"] for r in rooms if r["poly"] is not None]
    cay_p = shapely.STRtree(polys) if polys else None
    out, seen = [], set()
    ents = list(tpp.walk(msp))
    # net canh cua (doan thang tren layer cua): canh dang mo = 2 net song song chay doc tu ban le ra dau mo
    canh_l = []
    for e in ents:
        if e.dxftype() == "LINE" and tpp.layer_goc(e.dxf.layer) in want:
            a, b = (e.dxf.start.x, e.dxf.start.y), (e.dxf.end.x, e.dxf.end.y)
            if math.dist(a, b) > 300:
                canh_l.append(LineString([a, b]))
    cay_canh = shapely.STRtree(canh_l) if canh_l else None

    for e in ents:
        if e.dxftype() != "ARC" or tpp.layer_goc(e.dxf.layer) not in want_cung or not 500 <= e.dxf.radius <= 1200:
            continue
        try:
            h = e.ocs().to_wcs(e.dxf.center)
            p1, p2 = e.start_point, e.end_point
        except Exception:
            continue
        H, R = (h.x, h.y), e.dxf.radius
        k = (round(H[0] / 5), round(H[1] / 5), round(R))
        if k in seen:
            continue
        seen.add(k)
        ends = [(p1.x, p1.y), (p2.x, p2.y)]
        # dau 'dong' = dau ma doan H-E nam trong o cua (giua hai mat tuong, ngoai moi phong); canh mo nam trong phong.
        # Hoa thi xet them do song song voi canh tuong gan ban le (ban le sat goc tuong: ca hai dau deu song song).
        diem_j = []
        for j, E in enumerate(ends):
            mid_e = Point((H[0] + E[0]) / 2, (H[1] + E[1]) / 2)
            ngoai = 0 if (cay_p is not None and any(polys[i].contains(mid_e) for i in cay_p.query(mid_e))) else 1
            ss = 0.0
            if tree is not None:
                v = ((E[0] - H[0]) / R, (E[1] - H[1]) / R)
                for i in tree.query(Point(H).buffer(450)):
                    a, b = edges[i]
                    L = math.dist(a, b)
                    ss = max(ss, abs(v[0] * (b[0] - a[0]) / L + v[1] * (b[1] - a[1]) / L))
            diem_j.append((ngoai, round(ss, 3), -j))
        j = -max(diem_j)[2]
        nc = [dem_canh_cua(canh_l, cay_canh, H, E_, R) for E_ in ends]
        if nc[0] >= 2 > nc[1]:          # canh cua ve o dau 0 -> dau 0 la dau mo
            j = 1
        elif nc[1] >= 2 > nc[0]:
            j = 0
        E, M = ends[j], ends[1 - j]
        mid_ang = math.radians((e.dxf.start_angle + ((e.dxf.end_angle - e.dxf.start_angle) % 360) / 2))
        try:
            mc = e.ocs().to_wcs((e.dxf.center.x + R * 0.6 * math.cos(mid_ang), e.dxf.center.y + R * 0.6 * math.sin(mid_ang), 0))
            mid = (mc.x, mc.y)
        except Exception:
            mid = ((H[0] + (E[0] + M[0]) / 2) / 2, (H[1] + (E[1] + M[1]) / 2) / 2)
        out.append(dict(H=H, R=R, E=E, M=M, mid=mid, mo=LineString([H, E])))
    return out


def doc_hinh_lop(msp, layers):
    """Hinh hoc (LineString) cua cac net tren layer cho truoc (ke ca trong block)."""
    ch = tpp.doc_chains(msp, layers)
    return [LineString(p + ([p[0]] if c else [])) for c, p in ch if len(p) > 1]


def ap_nhan_dien_bo_sung(cfg, path):
    """File JSON {"<ten block>": "<loai noi that>"} (vd {"86786": "bep_nau", "CT3-Tn-MEP": "binh_nong_lanh"})."""
    if not path:
        return {}
    m = json.load(open(path, encoding="utf-8"))
    for ten, loai in m.items():
        if ten.startswith("_"):
            continue
        cfg["noi_that"].setdefault(loai, {"ten": []})
        pat = bo_dau(ten).replace("[", "[[]")
        cfg["noi_that"][loai]["ten"] = [pat] + cfg["noi_that"][loai].get("ten", [])
    return m


def doc_khoi_chua_ro(msp, doc, cfg, vung, da_nhan):
    """Block noi that CHUA nhan dien (de nguoi dung gan loai): tren layer noi that, kich thuoc 150..4500 mm."""
    lay = {tpp.layer_goc(x) for x in cfg.get("layer_noi_that_suy", [])}
    out = []

    def duyet(ents, depth):
        for e in ents:
            if e.dxftype() != "INSERT":
                continue
            bb = st.bb_ins(e, doc)
            if bb is None or not bb.intersects(vung):
                continue
            w, h = bb.bounds[2] - bb.bounds[0], bb.bounds[3] - bb.bounds[1]
            if max(w, h) > 4500:
                if depth < 3:
                    try:
                        duyet(e.virtual_entities(), depth + 1)
                    except Exception:
                        pass
                continue
            if tpp.layer_goc(e.dxf.layer) not in lay or max(w, h) < 150 or not vung.contains(bb.centroid):
                continue
            if any(abs(f["fp"].bounds[0] - bb.bounds[0]) < 2 and abs(f["fp"].bounds[1] - bb.bounds[1]) < 2 for f in da_nhan) or \
                    any(f["fp"].buffer(5).contains(bb) for f in da_nhan if f["loai"] in ("sofa", "ban_an", "giuong")):
                continue
            out.append(dict(ten=re.split(r"\$\d+\$", st.ten_that(e, doc))[-1], fp=bb, x=bb.centroid.x, y=bb.centroid.y,
                            rong=w, cao=h, layer=e.dxf.layer))

    duyet(msp, 0)
    # gop trung (cung ten + cung vi tri)
    gon, seen = [], set()
    for k in out:
        key = (k["ten"], round(k["x"]), round(k["y"]))
        if key not in seen:
            seen.add(key)
            gon.append(k)
    return gon


def doc_o_cam_co_san(msp, doc, catalog, vung):
    """Thiet bi dien da ve (SOAT): theo ten block (bi danh) + gia tri thuoc tinh -> ma thu vien."""
    bd = catalog["bi_danh_block"]
    by_block = defaultdict(list)
    for c in catalog["thiet_bi"]:
        by_block[c["block"]].append(c)

    def nhom(name):
        n = bo_dau(re.split(r"\$\d+\$", name)[-1])
        for blk, pats in bd.items():
            if blk.startswith("_"):
                continue
            if any(fnmatch.fnmatchcase(n, bo_dau(p)) for p in pats):
                return blk
        return None

    out = []

    def duyet(ents, depth):
        for e in ents:
            if e.dxftype() != "INSERT":
                continue
            name = st.ten_that(e, doc)
            blk = nhom(name)
            if blk is None:
                if depth < 2:
                    bb = st.bb_ins(e, doc)
                    if bb is not None and max(bb.bounds[2] - bb.bounds[0], bb.bounds[3] - bb.bounds[1]) > 3000 and bb.intersects(vung):
                        try:
                            duyet(e.virtual_entities(), depth + 1)
                        except Exception:
                            pass
                continue
            p = e.dxf.insert
            if not vung.contains(Point(p.x, p.y)):
                continue
            val = None
            try:
                for a in e.attribs:
                    val = (a.dxf.text or "").strip()
                    break
            except Exception:
                pass
            cands = by_block[blk]
            c = next((c for c in cands if (c["gia_tri"] or "") == (val or "")), None)
            if c is None:
                c = next((c for c in cands if bo_dau(c["gia_tri"] or "") == bo_dau(val or "")), None)
            if c is None:
                c = cands[0]
            m = e.matrix44()
            yv = m.transform_direction((0, 1, 0))
            out.append(dict(ma=c["ma"], cat=c, x=p.x, y=p.y, rot=e.dxf.rotation, val=val, block=name, layer=e.dxf.layer,
                            handle=e.dxf.handle, huong=(yv.x, yv.y)))

    duyet(msp, 0)
    return out


def doc_dim_va_nhan(msp):
    dims, nhan = [], []
    for e in msp:
        t = e.dxftype()
        if t == "DIMENSION":
            for k in ("defpoint2", "defpoint3"):
                p = e.dxf.get(k)
                if p is not None:
                    dims.append((p.x, p.y))
        elif t in ("TEXT", "MTEXT"):
            s = e.dxf.text if t == "TEXT" else e.plain_text()
            m = re.match(r"^\s*([A-Za-z]+\d*)\s*/\s*T\s*[ĐDđd]\.?\s*CH", s or "")
            if m:
                nhan.append(m.group(1).upper())
    return dims, nhan


# ================================================================================================= bien phong
class Bien:
    """Chu vi phong (mat tuong hoan thien) lay mau moi 10 mm: diem, canh, phap tuyen vao phong, co cam / BTCT."""
    BUOC = 10.0

    def __init__(self, P):
        self.P = orient(P, 1.0)
        self.ring = LineString(self.P.exterior.coords)
        self.L = self.ring.length
        c = list(self.P.exterior.coords)
        self.canh = []
        s0 = 0.0
        for a, b in zip(c, c[1:]):
            d = math.dist(a, b)
            if d < 1e-6:
                continue
            u = ((b[0] - a[0]) / d, (b[1] - a[1]) / d)
            self.canh.append(dict(a=a, b=b, L=d, s0=s0, u=u, n=(-u[1], u[0])))
            s0 += d
        n = max(2, int(self.L / self.BUOC))
        self.s = [i * self.L / n for i in range(n)]
        self.pts = [self.ring.interpolate(s) for s in self.s]
        self.cam = [None] * n          # ly do cam (chuoi) hoac None
        self.btct = [False] * n
        self.chiem = []                # doan (s_dau, s_cuoi) da dat thiet bi (thiet bi nho: s_dau = s_cuoi)

    def canh_tai(self, s):
        s %= self.L
        for k in self.canh:
            if k["s0"] - 1e-6 <= s <= k["s0"] + k["L"] + 1e-6:
                return k
        return self.canh[-1]

    def diem(self, s):
        s %= self.L
        p = self.ring.interpolate(s)
        k = self.canh_tai(s)
        return (p.x, p.y), k["n"], k

    def chieu(self, p):
        return self.ring.project(Point(p))

    def danh_dau(self, geom, d, ly_do):
        if geom is None or geom.is_empty:
            return
        g = prep(geom.buffer(d)) if d > 0 else prep(geom)
        for i, p in enumerate(self.pts):
            if self.cam[i] is None and g.contains(p):
                self.cam[i] = ly_do

    def danh_dau_btct(self, geom, d=40):
        if geom is None or geom.is_empty:
            return
        g = prep(geom.buffer(d))
        for i, p in enumerate(self.pts):
            if g.contains(p):
                self.btct[i] = True

    def ti_le_cam_canh(self, k, ly_do=("cửa / cửa sổ", "lan can / vách kính")):
        i0, i1 = self.idx(k["s0"] + 1), self.idx(k["s0"] + k["L"] - 1)
        ids = range(i0, i1 + 1) if i1 >= i0 else list(range(i0, len(self.s))) + list(range(0, i1 + 1))
        ids = list(ids)
        return sum(1 for i in ids if self.cam[i] in ly_do) / max(1, len(ids))

    def idx(self, s):
        return int(round((s % self.L) / self.L * len(self.s))) % len(self.s)

    def ly_do_cam(self, s):
        return self.cam[self.idx(s)]

    def la_btct(self, s):
        return self.btct[self.idx(s)]

    def trong(self, s, kc):
        def kc_doan(a, b):
            if a <= s <= b:
                return 0.0
            return min(min(abs(s - t), self.L - abs(s - t)) for t in (a, b))
        return all(kc_doan(a, b) >= kc for a, b in self.chiem)

    def tim(self, s0, dich_max, kc, tranh_btct=400, cung_canh=True, bo_qua=()):
        """Vi tri hop le gan s0 nhat (khong cam, cach thiet bi da dat >= kc); uu tien tuong xay trong pham vi tranh_btct.
        Tra ve (s, ghi_chu) hoac (None, ly do)."""
        k0 = self.canh_tai(s0)
        n = len(self.s)
        buoc = self.L / n
        ung = []
        for j in range(0, int(dich_max / buoc) + 1):
            for sg in ((1, -1) if j else (1,)):
                s = (s0 + sg * j * buoc) % self.L
                if cung_canh and self.canh_tai(s) is not k0 and j * buoc > 50:
                    continue
                if (self.ly_do_cam(s) is None or self.ly_do_cam(s) in bo_qua) and self.trong(s, kc):
                    ung.append((j * buoc, s))
            if ung and not tranh_btct:
                break
        if not ung:
            ld = self.ly_do_cam(s0) or "trùng thiết bị khác"
            return None, ld
        d0, s0b = min(ung)
        xay = [u for u in ung if not self.la_btct(u[1]) and u[0] <= max(tranh_btct, d0)]
        if xay:
            d, s = min(xay)
            if d > d0 + 15:
                return s, "dịch %d mm tránh vách BTCT" % d
            return s, ("dịch %d mm" % d) if d > 15 else ""
        return s0b, ("dịch %d mm, " % d0 if d0 > 15 else "") + "trên vách BTCT"


# ================================================================================================= bo tri
class BoTri:
    def __init__(self, cfg, catalog, args, so):
        self.cfg, self.b = cfg, cfg["bo_tri"]
        self.cat = {c["ma"]: c for c in catalog["thiet_bi"]}
        self.args, self.so = args, so
        self.ds = []            # thiet bi
        self.dims = []          # (p1, p2, p3)
        self.texts = []         # (x, y, chuoi, cao)
        self.may = []           # (x, y)

    def dat(self, ma, r, bien, s, ly_do, lech_vao=0.0, ghi_chu=""):
        (x, y), n, k = bien.diem(s)
        c = self.cat[ma]
        rot = (deg(n) - 90.0) % 360.0
        if ma == "VDP":
            rot = deg(n) % 360.0
        d = dict(ma=ma, cat=c, x=x + n[0] * lech_vao, y=y + n[1] * lech_vao, rot=rot, n=n, s=s, bien=bien, canh=k,
                 can=r["can"], phong=r["ten"], loai_phong=r["loai"], ly_do=ly_do, ghi_chu=ghi_chu, lo=None,
                 btct=bien.la_btct(s), co_dim=lech_vao == 0)
        # giu cho theo be rong ky hieu (tu dien ~500, man hinh VDP 250): thiet bi khac khong chen vao ky hieu
        nua = (c["rong"] / 2 - 60) if (c.get("rong") or 0) > 220 else 0.0
        d["_chiem"] = (s - nua, s + nua)
        bien.chiem.append(d["_chiem"])
        self.ds.append(d)
        if d["btct"]:
            self.may.append((x + n[0] * 120, y + n[1] * 120))
            self.so.them(r["can"], r["ten"], CB, "Thiết bị trên vách BTCT",
                         f"{c['ten']} không dời được sang tường xây trong phạm vi {self.b['tranh_btct_dich_toi_da']} mm: đặt chờ khi đổ cột vách (đã đánh dấu mây).", tb=d)
        return d

    def dat_gan(self, ma, r, bien, p_muc, ly_do, dich_max=600, kc=None, lech_vao=0.0, cung_canh=True, im_lang=False, bo_qua=()):
        """Dat thiet bi tai diem tuong gan p_muc nhat thoa rang buoc. im_lang: khong ghi chu khi that bai (con phuong an khac)."""
        kc = self.b["cach_nhau_toi_thieu"] if kc is None else kc
        s0 = bien.chieu(p_muc)
        s, gc = bien.tim(s0, dich_max, kc, self.b["tranh_btct_dich_toi_da"], cung_canh, bo_qua)
        if s is None and im_lang:
            return None
        if s is None:
            self.so.them(r["can"], r["ten"], CB, "Không đặt được thiết bị",
                         f"{self.cat[ma]['ten']}: vị trí theo nguyên tắc vướng '{gc}', không còn chỗ hợp lệ trong {dich_max} mm. Cần kiến trúc sư chọn vị trí.",
                         hoi=True, de_xuat=bien.diem(s0)[0])
            return None
        if gc and gc.startswith("dịch") and float(gc.split()[1]) > 150:
            self.so.them(r["can"], r["ten"], GY, "Dịch vị trí so với nguyên tắc", f"{self.cat[ma]['ten']}: {gc} (vướng {bien.ly_do_cam(s0) or 'vách BTCT / thiết bị khác'}).")
        return self.dat(ma, r, bien, s, ly_do, lech_vao, gc)


def tuong_gan(bien, fp):
    """Tuong noi that dua vao: trong cac canh gan nhat (<= dmin + 100) chon canh co doan chieu cua noi that dai nhat;
    bo canh ngan < 250 va canh phan lon la cua / cua so / lan can (khong lap thiet bi). Tra ve (kc, canh, t_min, t_max, t_tam)."""
    ung = []
    for k in bien.canh:
        if k["L"] < 250 or bien.ti_le_cam_canh(k) > 0.7:
            continue
        d = LineString([k["a"], k["b"]]).distance(fp)
        ts = [(x - k["a"][0]) * k["u"][0] + (y - k["a"][1]) * k["u"][1] for x, y in fp.exterior.coords]
        chong = max(0.0, min(max(ts), k["L"]) - max(min(ts), 0.0))
        ung.append((d, chong, k, min(ts), max(ts)))
    if not ung:
        ung = [(LineString([k["a"], k["b"]]).distance(fp), 0.0, k, 0.0, k["L"]) for k in bien.canh]
    dmin = min(u[0] for u in ung)
    d, _, k, t0, t1 = max((u for u in ung if u[0] <= dmin + 100), key=lambda u: (u[1], -u[0]))
    c = fp.centroid
    tc = (c.x - k["a"][0]) * k["u"][0] + (c.y - k["a"][1]) * k["u"][1]
    return d, k, t0, t1, tc


def diem_tren_canh(k, t):
    return (k["a"][0] + k["u"][0] * t, k["a"][1] + k["u"][1] * t)


def phong_cua(cuas, P, tol=300):
    """Cua cham chu vi phong P (o cua gan bien)."""
    return [c for c in cuas if c["mo"].distance(P.exterior) < tol]


def bo_tri_can(bt, can_ten, rooms, nt, cuas, can_poly, args, ke_tv_suy):
    """Bo tri mot can. rooms: phong cua can (co bien). Tra ve so phong ngu."""
    b, so = bt.b, bt.so
    kc = b["cach_nhau_toi_thieu"]

    def nt_trong(r, loai):
        return [f for f in nt if f["loai"] == loai and (r["poly"].buffer(150).contains(Point(f["x"], f["y"])))]

    so_pn = 0
    # ------------------------------------------------------------------ cua chinh: tu dien, VDP
    cua_chinh = None
    ung = []
    for c in cuas:
        H, E = c["H"], c["E"]
        mx, my = (H[0] + E[0]) / 2, (H[1] + E[1]) / 2
        L = math.dist(H, E) or 1
        nx, ny = -(E[1] - H[1]) / L, (E[0] - H[0]) / L
        a, bb_ = Point(mx + nx * 400, my + ny * 400), Point(mx - nx * 400, my - ny * 400)
        ia, ib = can_poly.buffer(-50).contains(a), can_poly.buffer(-50).contains(bb_)
        if ia == ib:
            continue
        ben_trong, ben_ngoai = (a, bb_) if ia else (bb_, a)
        r_in = next((r for r in rooms if r["poly"].contains(ben_trong)), None)
        if r_in is None or {"logia", "wc"} & set(r_in["loai"]):
            continue
        uu = 2 if {"khach", "hanh_lang"} & set(r_in["loai"]) else (1 if "bep" in r_in["loai"] else 0)
        ung.append((uu, c["R"], c, r_in, ben_trong))
    if ung:
        _, _, c0, r0, p0 = max(ung, key=lambda u: (u[0], u[1]))
        cua_chinh = (c0, r0, p0)
        if len(ung) > 1:
            so.them(can_ten, r0["ten"], GY, "Chọn cửa chính", f"{len(ung)} cửa đi nối căn với bên ngoài: chọn cửa rộng {c0['R']:.0f} mm vào {r0['ten']} làm cửa chính – kiểm tra.")
    if cua_chinh is None:
        so.them(can_ten, "", CB, "Không xác định được cửa chính", "Không thấy cửa đi có cung mở nối căn với bên ngoài: chưa bố trí TĐ-CH, VDP. Cần chỉ vị trí cửa chính.", hoi=True)
    else:
        c, r_in, _ = cua_chinh
        _cum_cua_vao(bt, can_ten, c, r_in)
    # ------------------------------------------------------------------ tung phong (WC truoc: cong tac ngoai cua WC gan voi
    # vi tri cua, uu tien hon o cam thuong cua phong ben ngoai)
    for r in sorted(rooms, key=lambda r: (0 if "wc" in r["loai"] else 1, r["ten"])):
        P, bien, loai = r["poly"], r["bien"], set(r["loai"])
        ten = r["ten"]
        cua_p = phong_cua(cuas, P)
        if loai & set(b["phong_khong_co_quy_tac"]) and not any(nt_trong(r, x) for x in ("giuong", "bep_nau", "may_giat", "ke_tv", "sofa")):
            so.them(can_ten, ten, GY, "Phòng chưa có quy tắc", f"{ten}: chưa có nguyên tắc bố trí ổ cắm (phòng đa năng) – hỏi người dùng nếu cần.", hoi=True)
        # ---------------- phong ngu
        if "ngu" in loai:
            so_pn += 1
            gs = nt_trong(r, "giuong")
            if not gs:
                so.them(can_ten, ten, LOI, "Thiếu nội thất", "Phòng ngủ không nhận diện được giường: chưa bố trí ổ G / TV. Gán loại block giường (--nhan-dien-bo-sung) hoặc chỉ vị trí.", hoi=True)
            for g in gs[:1]:
                than = g.get("than") or g["fp"]
                goi = g.get("vung_goi")
                d, k, t0, t1, tc = tuong_gan(bien, goi if goi is not None else than)
                if d > 400:
                    so.them(can_ten, ten, CB, "Giường không sát tường", f"Đầu giường cách tường {d:.0f} mm: ổ G đặt trên tường đầu giường gần nhất, cần kiểm tra.", hoi=True)
                _, _, b0, b1, _ = tuong_gan_canh(k, than)
                if g.get("truc") is None:
                    so.them(can_ten, ten, GY, "Giường không có tab đầu giường", "Mép giường lấy theo hộp bao block (có thể gồm tab đầu giường): kiểm tra vị trí ổ G.")
                gcach = b["g_cach_mep_giuong"]
                # vat can hai ben dau giuong tren tuong dau giuong: goc phong (dau canh) + noi that sat tuong (tu, ban...)
                vc_trai, vc_phai = [0.0], [k["L"]]
                for f in nt:
                    if f["loai"] in ("giuong", "ke_tv") or not r["poly"].buffer(150).contains(Point(f["x"], f["y"])):
                        continue
                    if LineString([k["a"], k["b"]]).distance(f["fp"]) > b["g_noi_that_sat_tuong"]:
                        continue
                    _, _, a0, a1, _ = tuong_gan_canh(k, f["fp"])
                    if a1 <= b0 + 1:
                        vc_trai.append(a1)
                    elif a0 >= b1 - 1:
                        vc_phai.append(a0)
                for mep, vc, sg in ((b0, max(vc_trai), -1), (b1, min(vc_phai), 1)):
                    khe = abs(mep - vc)
                    if khe < b["g_khe_can_giua"]:
                        t, ly = (mep + vc) / 2, f"Ổ đầu giường: khe cạnh giường – tường / nội thất {khe:.0f} < {b['g_khe_can_giua']}: cân giữa khe, H+650"
                    else:
                        t, ly = mep + sg * gcach, f"Ổ đầu giường: tâm cách mép giường {gcach}, H+650"
                    d_g = bt.dat_gan("O-G", r, bien, diem_tren_canh(k, t), ly, dich_max=150, kc=80, im_lang=True)
                    if d_g is None:
                        # phia ngoai vuong cua so / tu: dat vao trong mep giuong (sau dau giuong) - nhu can mau .03
                        d_g = bt.dat_gan("O-G", r, bien, diem_tren_canh(k, mep - sg * gcach),
                                         f"Ổ đầu giường: phía ngoài mép giường vướng cửa sổ / thiết bị – đặt vào trong mép giường {gcach} (sau đầu giường), H+650",
                                         dich_max=150)
                        if d_g is not None:
                            so.them(can_ten, ten, GY, "Ổ G đặt trong mép giường",
                                    f"Phía ngoài mép giường vướng cửa sổ / thiết bị: ổ G đặt vào trong mép giường {gcach} (sau đầu giường) – kiểm tra.", tb=d_g)
                # TV
                tvs = nt_trong(r, "ke_tv") + [f for f in ke_tv_suy if r["poly"].contains(Point(f["x"], f["y"]))]
                if tvs:
                    tv = tvs[0]
                    _, k2, _, _, tc2 = tuong_gan(bien, tv["fp"])
                    _dat_tv(bt, r, bien, k2, tc2, args, "Ổ tivi + ổ mạng phòng ngủ: cặp hộp theo trục tivi, H+1100" +
                            (" (tivi suy theo hình dạng – xác nhận)" if tv.get("suy") else ""))
                elif args.tv_pn_theo_truc_giuong:
                    # tuong doi dien dau giuong, theo truc giuong
                    cx, cy = than.centroid.x, than.centroid.y
                    nx, ny = k["n"]
                    ray = LineString([(cx, cy), (cx + nx * 20000, cy + ny * 20000)])
                    hit = ray.intersection(P.exterior)
                    pts = [hit] if hit.geom_type == "Point" else list(getattr(hit, "geoms", []))
                    if pts:
                        q = min(pts, key=lambda p: p.distance(Point(cx, cy)))
                        bt.dat_gan("O-TV", r, bien, (q.x, q.y), "Ổ tivi: theo trục giường, tường đối diện đầu giường (không có kệ TV – người dùng chọn)")
                else:
                    so.them(can_ten, ten, CB, "Thiếu kệ TV", "Phòng ngủ không có kệ TV / tivi: chưa đặt ổ TV. Đặt theo trục giường (--tv-pn-theo-truc-giuong) hay bỏ?", hoi=True)
            # o doi thuong gan cua phong
            if cua_p:
                c = cua_phong(cua_p, r, rooms)
                sH, sE = bien.chieu(c["H"]), bien.chieu(c["E"])
                huong = 1 if ((sE - sH) % bien.L) < bien.L / 2 else -1
                ly = f"Ổ đôi thường: gần cửa phòng, phía tay nắm, cách khuôn {b['o_thuong_pn_cach_khuon_cua']} mm"
                d0 = bt.dat_gan("O-DOI", r, bien, bien.diem(sE + huong * b["o_thuong_pn_cach_khuon_cua"])[0], ly, dich_max=900, im_lang=True)
                if d0 is None:      # phia tay nam vuong tu ao / cua khac: phia ban le (ngoai vung canh quet), roi tuong ke
                    d0 = bt.dat_gan("O-DOI", r, bien, bien.diem(sH - huong * (c["R"] + 100))[0],
                                    "Ổ đôi thường: gần cửa phòng, phía bản lề ngoài vùng cánh quét (phía tay nắm vướng)", dich_max=600, im_lang=True)
                if d0 is None:
                    bt.dat_gan("O-DOI", r, bien, bien.diem(sE + huong * b["o_thuong_pn_cach_khuon_cua"])[0],
                               "Ổ đôi thường: gần cửa phòng (dời sang tường kề, phía tay nắm vướng)", dich_max=1500, cung_canh=False)
            else:
                so.them(can_ten, ten, CB, "Không thấy cửa phòng", "Chưa đặt ổ đôi thường gần cửa (không nhận được cung mở cửa).", hoi=True)
            for f in nt_trong(r, "ban_lam_viec"):
                _, k3, _, _, tc3 = tuong_gan(bien, f["fp"])
                bt.dat_gan("O-DOI", r, bien, diem_tren_canh(k3, tc3), "Ổ bàn làm việc: theo tâm bàn (cao độ chưa chốt)")
        # ---------------- phong khach: TV + DN, sofa
        if "khach" in loai:
            tvs = nt_trong(r, "ke_tv") + [f for f in ke_tv_suy if r["poly"].contains(Point(f["x"], f["y"]))]
            sofa = nt_trong(r, "sofa")
            if tvs:
                tv = max(tvs, key=lambda f: f["fp"].area)
                _, k2, _, _, tc2 = tuong_gan(bien, tv["fp"])
                d1 = _dat_tv(bt, r, bien, k2, tc2, args, "Ổ tivi + ổ mạng phòng khách: 2 bộ theo trục tivi, H+1100 và H+400")
                if d1 is not None:
                    d1["ghi_chu"] = "2 bộ: H+1100 & H+400"
                    nx, ny = d1["n"]
                    bt.texts.append((d1["x"] + nx * 450 - ny * 150, d1["y"] + ny * 450 + nx * 150, "2 BỘ H+1100 & H+400", 80))
            else:
                so.them(can_ten, ten, CB, "Thiếu kệ TV", "Phòng khách không nhận diện được kệ TV / tivi: chưa đặt ổ TV + ĐN. Gán loại block hoặc chỉ vị trí.", hoi=True)
            for sf in sofa[:1]:
                d, k, t0, t1, tc = tuong_gan(bien, sf["fp"])
                if d > 500:
                    so.them(can_ten, ten, CB, "Sofa không sát tường", f"Sofa cách tường {d:.0f} mm: chưa đặt 2 ổ hai đầu sofa – hỏi vị trí.", hoi=True)
                    continue
                for t in (t0 - b["sofa_cach_mep"], t1 + b["sofa_cach_mep"]):
                    bt.dat_gan("O-DOI", r, bien, diem_tren_canh(k, t), f"Ổ đầu sofa: cách mép sofa {b['sofa_cach_mep']} mm", dich_max=300)
                if max(sf["fp"].bounds[2] - sf["fp"].bounds[0], sf["fp"].bounds[3] - sf["fp"].bounds[1]) > 3000:
                    so.them(can_ten, ten, GY, "Bộ sofa dạng cụm", "Block sofa gồm cả bàn/ghế phụ: mép sofa lấy theo hộp bao cụm – kiểm tra vị trí 2 ổ đầu sofa.")
        # ---------------- bep (theo noi that, phong nao co bep nau)
        bep = nt_trong(r, "bep_nau")
        if bep:
            h = bep[0]
            d, k, t0, t1, tc = tuong_gan(bien, h["fp"])
            q = diem_tren_canh(k, tc)
            dh = bt.dat_gan("HOP-HM", r, bien, q, "Box chờ hút mùi: theo tâm máy hút mùi / bếp nấu, H+1900", dich_max=200)
            if dh is not None:
                bt.ds.append(dict(dh, ma="HOP-BT", cat=bt.cat["HOP-BT"], x=dh["x"] + dh["n"][0] * b["bt_lech_vao_phong"],
                                  y=dh["y"] + dh["n"][1] * b["bt_lech_vao_phong"], co_dim=False, lo=None,
                                  ly_do="Đầu chờ bếp + ổ lò vi sóng: dưới bếp nấu, theo tâm bếp, H+500 (ký hiệu lệch 200 vào phòng để không chồng HM)"))
            # o B tren mat bep: 2 dau mat bep, tam cach dau mat bep 200, H+1300; nhieu hon 2 -> chia deu phan giua
            n_b = so_o_bep(args, can_ten)
            chau = [f for f in nt_trong(r, "chau_rua")]
            mb = []
            for f in nt_trong(r, "mat_bep"):          # mep truoc mat bep: song song tuong bep nau, cach tuong <= 900
                if "p0" not in f:
                    continue
                (x0, y0), (x1, y1) = f["p0"], f["p1"]
                dl = math.dist((x0, y0), (x1, y1)) or 1.0
                if abs((x1 - x0) / dl * k["u"][0] + (y1 - y0) / dl * k["u"][1]) > 0.98 and \
                        LineString([k["a"], k["b"]]).distance(LineString([f["p0"], f["p1"]])) <= 900:
                    mb.append(f)
            if mb:
                _, _, c0, c1, _ = tuong_gan_canh(k, max(mb, key=lambda f: f["fp"].area)["fp"])
            else:           # khong nhan duoc mat bep: theo thiet bi bep tren cung tuong +- 600
                ts = []
                for f in [h] + [f for f in chau + nt_trong(r, "may_rua_bat") + nt_trong(r, "lo_nuong")
                                if tuong_gan(bien, f["fp"])[1] is k and tuong_gan(bien, f["fp"])[0] < 900]:
                    _, _, a0, a1, _ = tuong_gan_canh(k, f["fp"])
                    ts += [a0, a1]
                c0, c1 = min(ts) - 600, max(ts) + 600
                so.them(can_ten, ten, GY, "Không nhận được mặt bếp", "Đầu mặt bếp lấy theo thiết bị bếp ± 600 – kiểm tra vị trí ổ B.")
            for f in nt_trong(r, "tu_lanh"):         # tu lanh dung canh mat bep cat dai mat bep
                _, k4, a0, a1, _ = tuong_gan(bien, f["fp"])
                if k4 is k:
                    if a1 <= tc:
                        c0 = max(c0, a1)
                    elif a0 >= tc:
                        c1 = min(c1, a0)
            c0, c1 = max(c0, 0.0), min(c1, k["L"])
            if n_b is None:
                so.them(can_ten, ten, CB, "Hỏi số ổ cắm mặt bếp",
                        f"Mặt bếp dài {(c1 - c0) / 1000:.1f} m (chi tiết lắp đặt: ổ ở 2 đầu mặt bếp). Hỏi số lượng trước khi bố trí (--so-o-bep).", hoi=True)
            else:
                dc = b["b_cach_dau_mat_bep"]
                vt = [c0 + dc, c1 - dc]
                if n_b == 1:
                    vt = [max(vt, key=lambda t: abs(t - tc))]
                elif n_b > 2:
                    cam_t = [(t0 - b["b_cach_mep_bep_nau"], t1 + b["b_cach_mep_bep_nau"])]
                    for f in chau:
                        _, k4, a0, a1, _ = tuong_gan(bien, f["fp"])
                        if k4 is k:
                            cam_t.append((a0 - b["b_cach_mep_chau_rua"], a1 + b["b_cach_mep_chau_rua"]))
                    vt += _chia_deu(_tru_khoang([(c0 + dc + 300, c1 - dc - 300)], cam_t), n_b - 2)
                for t in vt:
                    bt.dat_gan("O-B", r, bien, diem_tren_canh(k, t), f"Ổ cắm mặt bếp: tâm cách đầu mặt bếp {dc}, H+1300 (số lượng người dùng xác nhận)",
                               dich_max=300)
            for f in nt_trong(r, "may_rua_bat"):
                if args.may_rua_bat:
                    _, k5, _, _, tc5 = tuong_gan(bien, f["fp"])
                    bt.dat_gan("O-DOI", r, bien, diem_tren_canh(k5, tc5), "Ổ máy rửa bát (người dùng xác nhận; ký hiệu / cao độ chưa chốt)")
                else:
                    so.them(can_ten, ten, CB, "Hỏi ổ máy rửa bát", "Có block máy rửa bát: bố trí ổ cắm không? (--may-rua-bat)", hoi=True)
            for f in nt_trong(r, "lo_nuong"):
                if args.lo_nuong:
                    _, k5, _, _, tc5 = tuong_gan(bien, f["fp"])
                    bt.dat_gan("O-DOI", r, bien, diem_tren_canh(k5, tc5), "Ổ lò nướng (người dùng xác nhận; ký hiệu / cao độ chưa chốt)")
                else:
                    so.them(can_ten, ten, CB, "Hỏi ổ lò nướng", "Có block lò nướng / lò vi sóng: bố trí ổ cắm không? (--lo-nuong)", hoi=True)
        elif "bep" in loai:
            so.them(can_ten, ten, LOI, "Thiếu nội thất", "Bếp không nhận diện được bếp nấu: chưa bố trí BT / HM / ổ B. Gán loại block bếp nấu.", hoi=True)
        for f in nt_trong(r, "tu_lanh"):
            _, k6, _, _, tc6 = tuong_gan(bien, f["fp"])
            bt.dat_gan("O-TL", r, bien, diem_tren_canh(k6, tc6), "Ổ tủ lạnh: sau tủ lạnh, cân giữa tủ, H+1300")
        if (bep or "bep" in loai) and not nt_trong(r, "tu_lanh"):
            so.them(can_ten, ten, CB, "Thiếu tủ lạnh", "Khu bếp không nhận diện được tủ lạnh: chưa đặt ổ TL. Gán loại block hoặc chỉ vị trí.", hoi=True)
        # ---------------- WC
        if "wc" in loai:
            lav = nt_trong(r, "chau_rua")
            bc = nt_trong(r, "bon_cau")
            sen = nt_trong(r, "sen_tam") + nt_trong(r, "vach_tam")
            ref_sen = unary_union([f["fp"] for f in sen]) if sen else None
            if lav:
                _, k7, a0, a1, tc7 = tuong_gan(bien, lav[0]["fp"])
                cw = b["w_cach_mep_lavabo"]
                cands = [a0 - cw, a1 + cw]
                if ref_sen is not None:
                    cands.sort(key=lambda t: -Point(diem_tren_canh(k7, t)).distance(ref_sen))
                d_w = None
                for t in [t for t in cands if 60 <= t <= k7["L"] - 60]:      # lavabo sat goc: vi tri ngoai doan tuong -> bo
                    d_w = bt.dat_gan("O-W", r, bien, diem_tren_canh(k7, t), f"Ổ W chống ẩm lavabo: tường sau lavabo, tâm cách mép lavabo {cw}, H+1300",
                                     dich_max=100, im_lang=True)
                    if d_w is not None:
                        break
                if d_w is None:
                    # lavabo trong hoc / sat tuong ben: "mang tuong canh ngan lavabo" = tuong ben, tam cach mep truoc lavabo 300
                    sau_lav = max(abs((x - k7["a"][0]) * k7["n"][0] + (y - k7["a"][1]) * k7["n"][1]) for x, y in lav[0]["fp"].exterior.coords)
                    i7 = bien.canh.index(k7)
                    nk = len(bien.canh)
                    ben = []
                    for chieu in (-1, 1):
                        j = i7
                        for _ in range(6):          # bo qua doan song song (noi tiep) va khac ngan < 150
                            j = (j + chieu) % nk
                            ks = bien.canh[j]
                            if abs(ks["u"][0] * k7["u"][0] + ks["u"][1] * k7["u"][1]) <= 0.3 and ks["L"] >= 150:
                                break
                        else:
                            continue
                        if ks["L"] < sau_lav + cw + 60:
                            continue
                        goc = ks["b"] if chieu < 0 else ks["a"]          # dau canh tuong ben phia lavabo
                        hg = (-ks["u"][0], -ks["u"][1]) if chieu < 0 else ks["u"]
                        ben.append((goc[0] + hg[0] * (sau_lav + cw), goc[1] + hg[1] * (sau_lav + cw)))
                    if ref_sen is not None:
                        ben.sort(key=lambda p: -Point(p).distance(ref_sen))
                    for p in ben:
                        d_w = bt.dat_gan("O-W", r, bien, p, f"Ổ W chống ẩm lavabo: mảng tường cạnh ngắn lavabo (lavabo trong hốc), tâm cách mép trước lavabo {cw}, H+1300",
                                         dich_max=150, im_lang=True)
                        if d_w is not None:
                            break
                if d_w is None:
                    bt.dat_gan("O-W", r, bien, diem_tren_canh(k7, min(max(cands[0], 0.0), k7["L"])),
                               "Ổ W chống ẩm lavabo: gần lavabo (tường sau và tường bên vướng), H+1300", dich_max=600, cung_canh=False)
            else:
                so.them(can_ten, ten, CB, "Thiếu lavabo", "WC không nhận diện được lavabo: chưa đặt ổ W.", hoi=True)
            if bc:
                _, k8, a0, a1, tc8 = tuong_gan(bien, bc[0]["fp"])
                cx_ = b["x_cach_truc_bon_cau"]
                cands = [tc8 + cx_, tc8 - cx_]
                xa = (lav[0]["fp"] if lav else None)
                if xa is not None:
                    cands.sort(key=lambda t: -Point(diem_tren_canh(k8, t)).distance(xa))
                d_x = None
                for t in [t for t in cands if 60 <= t <= k8["L"] - 60]:
                    d_x = bt.dat_gan("O-X", r, bien, diem_tren_canh(k8, t), f"Ổ X chống ẩm bồn cầu: tường sau bồn cầu, tâm cách trục bồn cầu {cx_}, H+400",
                                     dich_max=100, im_lang=True)
                    if d_x is not None:
                        break
                if d_x is None:
                    bt.dat_gan("O-X", r, bien, diem_tren_canh(k8, min(max(cands[0], 0.0), k8["L"])),
                               "Ổ X chống ẩm bồn cầu (tường sau không đủ chỗ – dời sang tường kề), H+400", dich_max=500, cung_canh=False)
            else:
                so.them(can_ten, ten, CB, "Thiếu bồn cầu", "WC không nhận diện được bồn cầu: chưa đặt ổ X.", hoi=True)
            bnl = nt_trong(r, "binh_nong_lanh")
            if bnl:
                _, k9, _, _, tc9 = tuong_gan(bien, bnl[0]["fp"])
                bt.dat_gan("HOP-BNL", r, bien, diem_tren_canh(k9, tc9), "Box chờ bình nóng lạnh: theo tâm bình, trong WC")
            elif args.bnl_mac_dinh and cua_p and _bnl_mac_dinh(bt, r, min(cua_p, key=lambda c: c["mo"].distance(P.exterior))):
                bnl = True
            else:
                so.them(can_ten, ten, CB, "Thiếu bình nóng lạnh", "WC không nhận diện được bình nóng lạnh: chưa đặt box BNL / công tắc 20A. Chỉ vị trí bình hoặc gán loại block.", hoi=True)
            # cong tac ngoai cua WC (phia tay nam)
            if cua_p:
                c = min(cua_p, key=lambda c: c["mo"].distance(P.exterior))
                ngoai = _phong_ben_kia(c, r, rooms)
                if ngoai is None:
                    so.them(can_ten, ten, CB, "Không xác định được phía ngoài cửa WC", "Chưa đặt công tắc 3 phím / 20A ngoài cửa WC.", hoi=True)
                else:
                    _cong_tac_wc(bt, can_ten, ten, c, ngoai, bool(bnl))
            else:
                so.them(can_ten, ten, CB, "Không thấy cửa WC", "Chưa đặt công tắc ngoài cửa WC.", hoi=True)
        # ---------------- may giat (lo gia / bat ky)
        for f in nt_trong(r, "may_giat"):
            _, k10, a0, a1, tc10 = tuong_gan(bien, f["fp"])
            cm = b["w_may_giat_cach_mep"]
            # tren may giat, tam cach mep may 200, phia gan dau canh (goc tuong) hon
            t = a0 + cm if a0 <= k10["L"] - a1 else a1 - cm
            bt.dat_gan("O-W", r, bien, diem_tren_canh(k10, t), f"Ổ W chống ẩm máy giặt: trên máy giặt, tâm cách mép máy {cm}, H+1300", dich_max=200)
        for f in nt_trong(r, "dan_nong"):
            _, k11, _, _, tc11 = tuong_gan(bien, f["fp"])
            d = bt.dat_gan("HOP-AC", r, bien, diem_tren_canh(k11, tc11), "Box chờ điều hòa tại dàn nóng (phương án kiến trúc), cách trần 300")
            if d is not None:
                d["ghi_chu"] = (d["ghi_chu"] + "; " if d["ghi_chu"] else "") + "cao độ: cách trần 300"
    if not any(f["loai"] == "dan_nong" and can_poly.buffer(300).contains(Point(f["x"], f["y"])) for f in nt):
        lanh = [f for f in nt if f["loai"] == "dan_lanh" and can_poly.contains(Point(f["x"], f["y"]))]
        so.them(can_ten, "", CB, "Thiếu dàn nóng", ("Có %d dàn lạnh nhưng " % len(lanh) if lanh else "") +
                "không nhận diện được dàn nóng điều hòa: chưa đặt box AC. Chỉ vị trí dàn nóng theo phương án kiến trúc.", hoi=True)
    return so_pn


def _don_vi(a, b):
    L = math.dist(a, b) or 1.0
    return ((b[0] - a[0]) / L, (b[1] - a[1]) / L)


def _doan_toi_goc(bien, s_tu, h, w, toi_da=6000.0):
    """Tu vi tri s_tu (khuon cua) di theo chieu h (+1 / -1 theo chu vi) doc tuong cua den GOC TUONG THAT: bo qua cac canh
    song song tuong cua va cac khac ngan < 200 (khung cua, bac tuong); dung o canh dai khong song song. Tra ve (s_goc, dai)
    voi dai = khoang cach do theo phuong tuong cua."""
    k0 = bien.canh_tai(s_tu + h * 60)
    p0, _, _ = bien.diem(s_tu)
    n = len(bien.canh)
    i = bien.canh.index(k0)
    goc = (k0["b"] if h > 0 else k0["a"]) if abs(k0["u"][0] * w[0] + k0["u"][1] * w[1]) > 0.95 else p0
    di = 0.0
    for _ in range(n):
        i = (i + h) % n
        k = bien.canh[i]
        if abs(k["u"][0] * w[0] + k["u"][1] * w[1]) > 0.95:
            goc = k["b"] if h > 0 else k["a"]
        elif k["L"] >= 200:
            break
        di += k["L"]
        if di > toi_da:
            break
    dai = abs((goc[0] - p0[0]) * w[0] + (goc[1] - p0[1]) * w[1])
    return bien.chieu(goc), dai


def _go(bt, ds):
    for d in ds:
        bt.ds.remove(d)
        x = d.get("_chiem", (d["s"], d["s"]))
        if d.get("bien") is not None and x in d["bien"].chiem:
            d["bien"].chiem.remove(x)


def _cum_cua_vao(bt, can_ten, c, r_in):
    """Cua chinh (chi tiet lap dat thiet bi dien can ho dien hinh, nguoi dung chot 08/10/2026), mat trong tuong co cua chinh:
    - phia tay nam: man hinh chuong cua H+1300 + o cam H+400 (cung truc), tam cach goc tuong vuong goc 200;
    - phia ban le: tu dien TD-CH sat goc tuong (H+1300); doan tuong < 600 -> can giua doan;
    - ngoai cua: chuong / camera phia tay nam, tam cach lo mo 200, H+1400."""
    b, so = bt.b, bt.so
    bien = r_in["bien"]
    sH, sE = bien.chieu(c["H"]), bien.chieu(c["E"])
    huong = 1 if ((sE - sH) % bien.L) < bien.L / 2 else -1        # +huong: phia tay nam (xa ban le)
    # ---- man hinh chuong cua + o cam H400
    wd = _don_vi(c["H"], c["E"])
    s_goc_tn, dai_tn = _doan_toi_goc(bien, sE, huong, wd)
    cg = b["vdp_trong_cach_goc"]
    if dai_tn >= cg + b["o_cach_khuon_cua_toi_thieu"]:
        s_v, ly_v = s_goc_tn - huong * cg, f"Màn hình chuông cửa: mặt trong tường cửa chính, phía tay nắm, tâm cách góc tường vuông góc {cg}, H+1300"
    else:
        s_v, ly_v = sE + huong * dai_tn / 2, f"Màn hình chuông cửa: đoạn tường phía tay nắm ngắn ({dai_tn:.0f}) – cân giữa, H+1300"
    d_v = bt.dat_gan("VDP", r_in, bien, bien.diem(s_v)[0], ly_v, dich_max=300, kc=150, im_lang=True)
    if d_v is None:
        d_v = bt.dat_gan("VDP", r_in, bien, bien.diem(s_v)[0], "Màn hình chuông cửa: gần cửa chính (dời sang tường kề, tường cửa vướng), H+1300",
                         dich_max=1500, kc=150, cung_canh=False)
    if d_v is not None:
        bt.ds.append(dict(d_v, ma="O-DOI", cat=bt.cat["O-DOI"], rot=(deg(d_v["n"]) - 90.0) % 360.0, co_dim=False, lo=None,
                          x=d_v["x"] + d_v["n"][0] * b["ky_hieu_chong_lech_vao"], y=d_v["y"] + d_v["n"][1] * b["ky_hieu_chong_lech_vao"],
                          ly_do="Ổ cắm dưới màn hình chuông cửa, cùng trục, H+400 (ký hiệu lệch vào phòng để không chồng màn hình)", ghi_chu=""))
    # ---- tu dien: phia ban le, sat goc; doan < 600 can giua
    s_goc_bl, dai_bl = _doan_toi_goc(bien, sH, -huong, wd)
    bq = ("sau cánh cửa", "cách khuôn cửa < 200")
    d_td = None
    if dai_bl >= b["td_doan_can_giua"]:
        d_td = bt.dat_gan("TU-DIEN", r_in, bien, bien.diem(s_goc_bl + huong * b["td_cach_goc"])[0],
                          f"Tủ điện: mặt trong tường cửa chính, phía bản lề, sát góc tường (tâm cách góc {b['td_cach_goc']}), H+1300",
                          dich_max=200, kc=300, im_lang=True, bo_qua=bq)
    elif dai_bl >= b["td_doan_toi_thieu"]:
        d_td = bt.dat_gan("TU-DIEN", r_in, bien, bien.diem(sH - huong * dai_bl / 2)[0],
                          f"Tủ điện: đoạn tường phía bản lề {dai_bl:.0f} < {b['td_doan_can_giua']} – cân giữa, H+1300",
                          dich_max=100, kc=300, im_lang=True, bo_qua=bq)
    if d_td is None:            # phia ban le khong du cho: tuong cua phia tay nam, giua man hinh va khuon cua; roi tuong ke
        s_td = (d_v["s"] if d_v is not None else s_goc_tn) - huong * 600
        d_td = bt.dat_gan("TU-DIEN", r_in, bien, bien.diem(s_td)[0], "Tủ điện: phía bản lề không đủ chỗ – đặt phía tay nắm, cạnh màn hình chuông cửa, H+1300",
                          dich_max=300, kc=300, im_lang=True)
    if d_td is None:
        d_td = bt.dat_gan("TU-DIEN", r_in, bien, bien.diem(s_goc_bl)[0], "Tủ điện: gần cửa chính, dời sang tường kề (tường cửa vướng), H+1300",
                          dich_max=1500, kc=300, cung_canh=False, bo_qua=bq)
    if d_td is not None and not d_td["ly_do"].startswith("Tủ điện: mặt trong tường cửa chính"):
        so.them(can_ten, r_in["ten"], GY, "Tủ điện không sát góc phía bản lề", d_td["ly_do"] + " – kiểm tra phối hợp nội thất (tủ giày).", tb=d_td)
    # ---- chuong / camera ngoai cua: mat ngoai tuong hanh lang, phia tay nam, tam cach lo mo 200, H+1400
    (px, py), n, _ = bien.diem(sE + huong * b["vdp_ngoai_cach_khuon_cua"])
    t = b["tuong_day_mac_dinh"]
    dd = _day_tuong(bt.tuong_net, (px, py), (-n[0], -n[1]))
    if dd:
        t = dd
    vx, vy = px - n[0] * t, py - n[1] * t
    bt.ds.append(dict(ma="VDP", cat=bt.cat["VDP"], x=vx, y=vy, rot=deg((-n[0], -n[1])) % 360, n=(-n[0], -n[1]), s=None, bien=None,
                      canh=None, can=can_ten, phong="Ngoài cửa chính", loai_phong=["ngoai"], lo=None, btct=False, co_dim=False,
                      ly_do="Chuông cửa / camera ngoài cửa chính, phía tay nắm, tâm cách lỗ mở 200, H+1400", ghi_chu=f"tường dày {t:.0f}"))


def _dat_tv(bt, r, bien, k2, tc2, args, ly):
    """Cap o cam tivi + o mang theo truc tivi: hai hop 86x86 sat nhau, tam cach 100 (de vuong), o cam ben trai truc."""
    kc_tv = bt.b["tv_dn_khoang_cach"][args.de_o]
    d1 = bt.dat_gan("O-TV", r, bien, diem_tren_canh(k2, tc2 - kc_tv / 2), ly, kc=kc_tv - 5)
    if d1 is not None:
        bt.dat_gan("O-DN", r, bien, bien.diem(d1["s"] + kc_tv)[0], f"Ổ mạng cạnh ổ tivi, tâm cách {kc_tv} (hai hộp sát nhau)", kc=kc_tv - 5, dich_max=60)
    return d1


def _cong_tac_wc(bt, can_ten, ten, c, ngoai, co_bnl):
    """Cong tac ngoai cua WC (chi tiet lap dat): phia khong co ban le (tay nam), H+1300; cong tac may nuoc nong (20A) gan lo
    mo, tam cach lo mo >= 200 (toi thieu 80), cong tac den / guong / quat hut ke tiep, tam cach 100. Doan tuong tu lo mo den
    goc < 400 -> can giua; khong du cho 2 mat -> lap mat tren / mat duoi (1300 & 1200) cung vi tri."""
    b, so = bt.b, bt.so
    b2 = ngoai["bien"]
    sH, sE = b2.chieu(c["H"]), b2.chieu(c["E"])
    huong = 1 if ((sE - sH) % b2.L) < b2.L / 2 else -1
    dc, kc2 = b["ct_wc_cach_khuon_cua"], b["ct_wc_khoang_cach"]
    ma_ds = (["CT-20A"] if co_bnl else []) + ["CT-BA"]
    ten_ct = {"CT-20A": f"Công tắc máy nước nóng 20A ngoài cửa {ten}", "CT-BA": f"Công tắc đèn / gương / quạt hút ngoài cửa {ten}"}
    bq = ("cách khuôn cửa < 200",)

    def cap_ngang(s_khuon, h, cho_80, ghi):
        ds = []
        for i, ma in enumerate(ma_ds):
            d = bt.dat_gan(ma, ngoai, b2, b2.diem(s_khuon + h * (dc + i * kc2))[0], ten_ct[ma] + ghi + ", H+1300", dich_max=150 if cho_80 else 60,
                           kc=kc2 - 5, im_lang=True, bo_qua=bq if cho_80 else ())
            if d is None:
                _go(bt, ds)
                return None
            ds.append(d)
        return ds

    def chong(s_muc, ghi, **kw):
        d = bt.dat_gan(ma_ds[0], ngoai, b2, b2.diem(s_muc)[0], ten_ct[ma_ds[0]] + ghi + ", mặt trên H+1300", kc=60, im_lang=True, bo_qua=bq, **kw)
        if d is None:
            return None
        ds = [d]
        if len(ma_ds) > 1:
            d2 = bt.dat(ma_ds[1], ngoai, b2, d["s"], ten_ct[ma_ds[1]] + ghi + ", mặt dưới H+1200 (lắp chồng)")
            ds.append(d2)
            so.them(can_ten, ten, GY, "Công tắc WC lắp chồng", "Không đủ chỗ cho 2 mặt công tắc cạnh nhau: lắp mặt trên / mặt dưới (1300 & 1200).", tb=d)
        return ds

    ds = None
    for h, phia in ((huong, "phía tay nắm"), (-huong, "phía bản lề (phía tay nắm vướng)")):
        s_khuon = sE if h == huong else sH
        _, dai = _doan_toi_goc(b2, s_khuon, h, _don_vi(c["H"], c["E"]))
        if dai < b["ct_wc_doan_can_giua"]:
            ds = chong(s_khuon + h * dai / 2, f" ({phia}, đoạn tường {dai:.0f} < {b['ct_wc_doan_can_giua']} – cân giữa)", dich_max=100)
        else:
            ds = cap_ngang(s_khuon, h, False, f" ({phia})") or cap_ngang(s_khuon, h, True, f" ({phia}, tối thiểu 80 từ lỗ mở)")
        if ds:
            break
    if not ds:
        ds = chong(sE + huong * dc, " (dời sang tường kề, gần cửa vướng)", dich_max=1200, cung_canh=False)
    if not ds:
        so.them(can_ten, ten, CB, "Không đặt được công tắc WC", "Không còn chỗ hợp lệ cho công tắc ngoài cửa WC – cần kiến trúc sư chọn vị trí.", hoi=True)
        return
    for d in ds:
        d["wc"] = ten


def _bnl_mac_dinh(bt, r, c):
    """WC khong co block binh nong lanh (--bnl-mac-dinh): box BNL tren tuong vuong goc tuong cua, tai goc phong gan ban le
    nhat, cach goc bnl_mac_dinh_cach_goc (theo can mau nguoi dung ve: binh dat tren cua, box o goc phia ban le)."""
    bien = r["bien"]
    H, E = c["H"], c["E"]
    L = math.dist(H, E) or 1.0
    w = ((E[0] - H[0]) / L, (E[1] - H[1]) / L)               # phuong tuong co cua
    # canh tuong vuong goc tuong cua, co dau mut gan ban le nhat (goc phong phia ban le; bo khac o cua ngan < 400)
    ung = []
    for k in bien.canh:
        if k["L"] < 400 or abs(k["u"][0] * w[0] + k["u"][1] * w[1]) > 0.3:
            continue
        for dau, p in (("a", k["a"]), ("b", k["b"])):
            if math.dist(p, H) <= 700:
                ung.append((math.dist(p, H), dau, k))
    if not ung:
        return False
    _, dau, k = min(ung, key=lambda u: u[0])
    d = bt.b["bnl_mac_dinh_cach_goc"]
    t = d if dau == "a" else k["L"] - d
    s = k["s0"] + t
    dd = bt.dat("HOP-BNL", r, bien, s, f"Box chờ bình nóng lạnh: không có block bình – đặt góc tường phía bản lề cửa {r['ten']}, "
                f"cách góc {d} (theo căn mẫu người dùng vẽ, box cao trên cửa)")
    dd["ghi_chu"] = "vị trí mặc định – xác nhận vị trí bình"
    bt.so.them(r["can"], r["ten"], CB, "BNL đặt theo vị trí mặc định",
               "Không có block bình nóng lạnh: box BNL đặt góc tường phía bản lề cửa WC (theo căn mẫu) – xác nhận vị trí bình.", tb=dd)
    return True


def tuong_gan_canh(k, fp):
    ts = [(x - k["a"][0]) * k["u"][0] + (y - k["a"][1]) * k["u"][1] for x, y in fp.exterior.coords]
    c = fp.centroid
    return None, k, min(ts), max(ts), (c.x - k["a"][0]) * k["u"][0] + (c.y - k["a"][1]) * k["u"][1]


def _tru_khoang(ds, cam):
    out = list(ds)
    for c0, c1 in cam:
        moi = []
        for a, b in out:
            if c1 <= a or c0 >= b:
                moi.append((a, b))
                continue
            if c0 > a:
                moi.append((a, c0))
            if c1 < b:
                moi.append((c1, b))
        out = moi
    return [(a, b) for a, b in out if b - a > 150]


def _chia_deu(khoang, n):
    if not khoang:
        return []
    dem = [0] * len(khoang)
    for _ in range(n):
        i = max(range(len(khoang)), key=lambda i: (khoang[i][1] - khoang[i][0]) / (dem[i] + 1))
        dem[i] += 1
    out = []
    for (a, b), c in zip(khoang, dem):
        out += [a + (b - a) * (j + 1) / (c + 1) for j in range(c)]
    return out


def cua_phong(cua_p, r, rooms):
    """Cua vao phong: uu tien cua thong ra P. khach / hanh lang / bep (khong phai cua WC rieng, lo gia); on dinh giua cac lan chay."""
    def diem(c):
        q = _phong_ben_kia(c, r, rooms)
        lq = set(q["loai"]) if q else set()
        uu = 3 if lq & {"khach", "hanh_lang"} else 2 if lq & {"bep", "an"} else 0 if lq & {"wc", "logia"} else 1
        return (uu, round(c["R"]), round(c["H"][0]), round(c["H"][1]))
    return max(cua_p, key=diem)


def _phong_ben_kia(c, r, rooms):
    H, E = c["H"], c["E"]
    mx, my = (H[0] + E[0]) / 2, (H[1] + E[1]) / 2
    L = math.dist(H, E) or 1
    nx, ny = -(E[1] - H[1]) / L, (E[0] - H[0]) / L
    for sg in (1, -1):
        p = Point(mx + sg * nx * 450, my + sg * ny * 450)
        if r["poly"].contains(p):
            continue
        q = next((x for x in rooms if x is not r and x["poly"].contains(p)), None)
        if q is not None:
            return q
    return None


def _day_tuong(net, p, v, toi_da=450):
    """Khoang cach tu mat tuong trong (p) theo huong v toi net tuong ngoai xa nhat <= toi_da (be day tuong)."""
    if net is None:
        return None
    ray = LineString([p, (p[0] + v[0] * toi_da, p[1] + v[1] * toi_da)])
    ds = []
    for i in net[1].query(ray):
        g = net[0][i].intersection(ray)
        for q in ([g] if g.geom_type == "Point" else list(getattr(g, "geoms", []))):
            if q.geom_type == "Point":
                d = math.dist(p, (q.x, q.y))
                if d > 60:
                    ds.append(d)
    return max(ds) if ds else None


def so_o_bep(args, can):
    v = args.so_o_bep
    if v is None:
        return None
    v = str(v)
    if "=" not in v:
        return int(v)
    for part in v.split(","):
        k, _, n = part.partition("=")
        if bo_dau(k.strip()) == bo_dau(can):
            return int(n)
    return None


# ================================================================================================= lo, cap, dim
def chia_lo(bt, can_ten, rooms, so_pn):
    b = bt.b
    ds = [d for d in bt.ds if d["can"] == can_ten]
    td = next((d for d in ds if d["ma"] == "TU-DIEN"), None)
    pn = [r for r in rooms if "ngu" in r["loai"]]
    if td is not None:
        pn.sort(key=lambda r: r["poly"].distance(Point(td["x"], td["y"])))
    tach = len(pn) > b["pn_tach_lo_khi_so_pn_lon_hon"]
    nhom_pn = {id(r): (0 if (not tach or i < (len(pn) + 1) // 2) else 1) for i, r in enumerate(pn)}

    def wc_rieng(r):
        """WC co cua mo ra phong ngu -> thuoc lo phong ngu."""
        for d in ds:
            if d["ma"] == "CT-BA" and d.get("wc") == r["ten"]:
                return next((x for x in rooms if x["ten"] == d["phong"] and "ngu" in x["loai"]), None)
        return None

    phong = {r["ten"]: r for r in rooms}
    dem = Counter()
    for d in ds:
        r = phong.get(d["phong"])
        m = d["ma"]
        if m == "HOP-AC":
            dem["AC"] += 1
            d["lo"] = f"AC{dem['AC']}"
        elif m == "HOP-BT":
            d["lo"] = "BT"
        elif m in ("HOP-BNL",):
            dem["HW"] += 1
            d["lo"] = f"HW{dem['HW']}"
            d["hw_wc"] = d["phong"]
        elif m == "TU-DIEN":
            d["lo"] = None
        elif m == "CT-BA":
            d["lo"] = None          # mach chieu sang / quat hut: ghi F
        elif m == "CT-20A":
            d["lo"] = "HW?"
        elif m in ("O-B", "O-TL", "HOP-HM") or (m == "O-W" and r is not None and "wc" not in r["loai"]) or \
                (m == "O-DOI" and ("máy rửa bát" in d["ly_do"] or "lò nướng" in d["ly_do"])):
            d["lo"] = "S-BEP"
        elif r is not None and "ngu" in r["loai"]:
            d["lo"] = f"S-PN{nhom_pn.get(id(r), 0) + 1}"
        elif r is not None and "wc" in r["loai"]:
            rr = wc_rieng(r)
            d["lo"] = f"S-PN{nhom_pn.get(id(rr), 0) + 1}" if rr is not None else "S-KHACH"
        else:
            d["lo"] = "S-KHACH"
    for d in ds:
        if d["ma"] == "CT-20A":
            hw = next((x for x in ds if x["ma"] == "HOP-BNL" and x["phong"] == d.get("wc")), None)
            d["lo"] = hw["lo"] if hw else None
    # dat ten S1.. theo thu tu khach, bep, PN1, PN2
    ten_s = {}
    k = 0
    for g in ("S-KHACH", "S-BEP", "S-PN1", "S-PN2"):
        if any(d["lo"] == g for d in ds):
            k += 1
            ten_s[g] = f"S{k}"
    for d in ds:
        if d["lo"] in ten_s:
            d["lo"] = ten_s[d["lo"]]
    return tach


def ve_dim(bt, can_ten):
    """Kich thuoc tu mep tuong (dau canh / mep o cua) den tam thiet bi; thiet bi lien tiep -> chuoi kich thuoc."""
    off = bt.b["dim_cach_tuong"]
    theo_canh = defaultdict(list)
    for d in bt.ds:
        if d["can"] != can_ten or not d.get("co_dim") or d["bien"] is None:
            continue
        theo_canh[(id(d["bien"]), id(d["canh"]))].append(d)
    for (_, _), mem in theo_canh.items():
        bien, k = mem[0]["bien"], mem[0]["canh"]
        t_of = lambda d: (d["x"] - k["a"][0]) * k["u"][0] + (d["y"] - k["a"][1]) * k["u"][1]    # noqa: E731
        # moc mep tuong: 2 dau canh + mep cac o cua tren canh
        refs = [0.0, k["L"]]
        i0, i1 = bien.idx(k["s0"]), bien.idx(k["s0"] + k["L"] - 1)
        prev = None
        for i in range(i0, i1 + 1):
            c = bien.cam[i] if bien.cam[i] in ("cửa / cửa sổ",) else None
            if (c is None) != (prev is None) and i > i0:
                refs.append(bien.s[i] - k["s0"])
            prev = c
        refs = sorted(set(round(r, 1) for r in refs if -1 <= r <= k["L"] + 1))
        mem = sorted(mem, key=t_of)
        ts = [t_of(d) for d in mem]
        nx, ny = k["n"]

        def P(t):
            return diem_tren_canh(k, t)

        def them(t1, t2):
            if abs(t2 - t1) < 20:
                return
            p1, p2 = P(t1), P(t2)
            p3 = (p1[0] + nx * off, p1[1] + ny * off)
            bt.dims.append((p1, p2, p3))
        # chia theo khoang giua hai moc lien tiep
        for rL, rR in zip(refs, refs[1:]):
            grp = [t for t in ts if rL - 1 <= t <= rR + 1]
            if not grp:
                continue
            giua = (rL + rR) / 2
            trai = [t for t in grp if t <= giua]
            phai = [t for t in grp if t > giua]
            last = rL
            for t in trai:
                them(last, t)
                last = t
            last = rR
            for t in reversed(phai):
                them(t, last)
                last = t


# ================================================================================================= dau ra
def xuat_scr(path, bt, cfg, catalog, chen_bang, vung_bang):
    v = cfg["ve"]
    L_ = v["layer"]

    def s(x):
        return '"' + tpp.acad_str(str(x)).replace("\\", "\\\\").replace('"', "'") + '"'

    def pt(p):
        return f"(list {p[0]:.1f} {p[1]:.1f} 0.0)"

    lay = {k: tpp.acad_str(x[0]) for k, x in L_.items()}
    out = ["(setvar \"CMDECHO\" 0)",
           "(setq _osm (getvar \"OSMODE\") _ar (getvar \"ATTREQ\") _ad (getvar \"ATTDIA\") _cl (getvar \"CLAYER\") _ab (getvar \"ANGBASE\") _adr (getvar \"ANGDIR\"))",
           "(setvar \"OSMODE\" 0)", "(setvar \"ATTREQ\" 1)", "(setvar \"ATTDIA\" 0)", "(setvar \"FILEDIA\" 0)", "(setvar \"ANGBASE\" 0.0)", "(setvar \"ANGDIR\" 0)",
           "(if (not (tblsearch \"LTYPE\" \"CENTER2\")) (command \"_.-LINETYPE\" \"_L\" \"CENTER2\" \"acadiso.lin\" \"\"))",
           "(defun mk-lay (n c lt) (if (not (tblsearch \"LAYER\" n)) (entmake (list (cons 0 \"LAYER\") (cons 100 \"AcDbSymbolTableRecord\") "
           "(cons 100 \"AcDbLayerTableRecord\") (cons 2 n) (cons 70 0) (cons 62 c) (cons 6 (if (tblsearch \"LTYPE\" lt) lt \"Continuous\"))))) n)"]
    for k, (n, c, lt) in L_.items():
        out.append(f"(mk-lay {s(n)} {c} {s(lt)})")
    out += [f"(if (not (tblsearch \"STYLE\" {s(v['text_style'])})) (entmake (list (cons 0 \"STYLE\") (cons 100 \"AcDbSymbolTableRecord\") "
            f"(cons 100 \"AcDbTextStyleTableRecord\") (cons 2 {s(v['text_style'])}) (cons 70 0) (cons 40 0.0) (cons 41 1.0) (cons 50 0.0) "
            f"(cons 71 0) (cons 42 100.0) (cons 3 {s(v['text_font'])}) (cons 4 \"\"))))"]
    # dim style
    db = v["dim_bien"]
    set_dim = " ".join(f"(setvar \"{k}\" {s(x) if isinstance(x, str) else x})" for k, x in db.items())
    out += [f"(if (tblsearch \"DIMSTYLE\" {s(v['dim_style'])}) (command \"_.-DIMSTYLE\" \"_R\" {s(v['dim_style'])}) "
            f"(progn {set_dim} (setvar \"DIMTXSTY\" {s(v['text_style'])}) (command \"_.-DIMSTYLE\" \"_S\" {s(v['dim_style'])})))"]
    # chen block: lan dau chua co dinh nghia -> "ten=file" (nap tu thu vien); KHONG dung "-INSERT ten=file" + (command) de huy
    # (Core Console thoat ngang)
    cat = {c["ma"]: c for c in catalog["thiet_bi"]}
    # Block DA CO trong ban ve co the mang don vi chen khac mm (vd block dong tu dien: inch) -> AutoCAD tu nhan 25,4 khi chen.
    # hsdv = (mm / don vi block) / (mm / don vi ban ve); chia ti le chen cho hsdv de ky hieu dung kich thuoc 1:1 mm.
    out += ["(defun mmdv (u) (cond ((= u 1) 25.4) ((= u 2) 304.8) ((= u 5) 10.0) ((= u 6) 1000.0) ((= u 14) 100.0) (T 1.0)))",
            "(defun hsdv (n / r u) (setq r (tblobjname \"BLOCK\" n)) (if r (setq u (cdr (assoc 70 (entget (cdr (assoc 330 (entget r)))))))) "
            "(if (and u (> u 0) (> (getvar \"INSUNITS\") 0)) (/ (mmdv u) (mmdv (getvar \"INSUNITS\"))) 1.0))"]

    def ten_blk(c):
        f = os.path.join(THU_VIEN, c["file"]).replace("\\", "/")
        return f"(if (tblsearch \"BLOCK\" {s(c['block'])}) {s(c['block'])} {s(c['block'] + '=' + f)})"

    def chen(c, x, y, rot, val, layer_key=None, sc=None):
        lk = layer_key or c["layer"]
        sc = sc or c.get("ty_le", 1.0)
        L = [f"(setvar \"CLAYER\" {s(L_[lk][0])})"]
        if c.get("tag"):
            L.append(f"(command \"_.-INSERT\" {ten_blk(c)} \"_S\" (/ {sc} (hsdv {s(c['block'])})) \"_R\" {rot:.3f} {pt((x, y))} {s(val or '')})")
        else:
            L.append(f"(command \"_.-INSERT\" {ten_blk(c)} \"_S\" (/ {sc} (hsdv {s(c['block'])})) \"_R\" {rot:.3f} {pt((x, y))})")
        return L

    for d in bt.ds:
        c = d["cat"]
        out += chen(c, d["x"], d["y"], d["rot"], c.get("gia_tri"))
        if c.get("nhan"):
            nx, ny = d["n"]
            out.append(f"(entmake (list (cons 0 \"TEXT\") (cons 8 {s(L_['text'][0])}) (cons 7 {s(v['text_style'])}) "
                       f"{'(list 10 %.1f %.1f 0.0)' % (d['x'] + nx * 300 - 150, d['y'] + ny * 300 + 120)} (cons 40 110.0) (cons 1 {s(c['nhan'])})))")
    for x, y in bt.may:
        out += chen(cat["MAY-CHO"], x, y, 0.0, None)
    for x, y, t, h in bt.texts:
        out.append(f"(entmake (list (cons 0 \"TEXT\") (cons 8 {s(L_['text'][0])}) (cons 7 {s(v['text_style'])}) (list 10 {x:.1f} {y:.1f} 0.0) "
                   f"(cons 40 {float(h):.1f}) (cons 1 {s(t)})))")
    out.append(f"(setvar \"CLAYER\" {s(L_['dim'][0])})")
    for p1, p2, p3 in bt.dims:
        out.append(f"(command \"_.DIMALIGNED\" {pt(p1)} {pt(p2)} {pt(p3)})")
    if chen_bang and vung_bang is not None:
        bk = catalog["bang_ky_hieu"]
        out.append(f"(setvar \"CLAYER\" {s(L_['text'][0])})")
        out.append(f"(command \"_.-INSERT\" {ten_blk(dict(block=bk['block'], file=bk['file']))} \"_S\" (/ 1.0 (hsdv {s(bk['block'])})) \"_R\" 0 {pt(vung_bang)})")
    out += ["(setvar \"CLAYER\" _cl)", "(setvar \"OSMODE\" _osm)", "(setvar \"ATTREQ\" _ar)", "(setvar \"ATTDIA\" _ad)",
            "(setvar \"ANGBASE\" _ab)", "(setvar \"ANGDIR\" _adr)", "_.QSAVE"]
    for line in out:
        if " " in THU_VIEN and "-INSERT" in line:
            raise SystemExit(f"Đường dẫn thư viện có dấu cách: {THU_VIEN}")
    with open(path, "w", encoding="ascii", errors="strict", newline="\n") as fh:
        fh.write("\n".join(out) + "\n")


def ve_anh(path, can_poly, rooms, nt, bt, cuas, so, ten_can, chua_ro, co_san=None, de_xuat=None):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    minx, miny, maxx, maxy = can_poly.bounds
    w, h = maxx - minx + 1, maxy - miny + 1
    fig = plt.figure(figsize=(14, 14 * h / w + 0.6), dpi=110)
    ax = fig.add_axes([0.01, 0.01, 0.98, 0.95])
    for r in rooms:
        xs, ys = r["poly"].exterior.xy
        ax.plot(xs, ys, color="#343a40", lw=1, ls="--" if r.get("gan_dung") else "-")
        c = r["poly"].representative_point()
        ax.text(c.x, c.y, r["ten"], fontsize=7, color="#1864ab", ha="center")
        bien = r.get("bien")
        if bien is not None:
            for i, p in enumerate(bien.pts):
                if i % 3:
                    continue
                if bien.cam[i]:
                    ax.plot(p.x, p.y, ".", color="#ffa8a8", ms=1.5)
                elif bien.btct[i]:
                    ax.plot(p.x, p.y, ".", color="#ced4da", ms=1.5)
    for f in nt:
        if not can_poly.buffer(300).contains(Point(f["x"], f["y"])):
            continue
        g = f["fp"]
        for gg in getattr(g, "geoms", [g]):
            xs, ys = gg.exterior.xy
            ax.plot(xs, ys, color="#adb5bd", lw=0.6)
        ax.text(f["x"], f["y"], f["loai"] + ("?" if f.get("suy") else ""), fontsize=5, color="#868e96", ha="center")
    for k in chua_ro:
        if can_poly.buffer(300).contains(Point(k["x"], k["y"])):
            xs, ys = k["fp"].exterior.xy
            ax.plot(xs, ys, color="#f08c00", lw=0.8, ls=":")
            ax.text(k["x"], k["y"], f"#{k['id']}", fontsize=6, color="#e8590c", ha="center")
    for c in cuas:
        if can_poly.buffer(300).contains(Point(c["H"])):
            ax.add_patch(matplotlib.patches.Circle(c["H"], c["R"], fill=False, ec="#ffc9c9", lw=0.5))
    mau = {}
    pal = ["#e03131", "#1971c2", "#2f9e44", "#f08c00", "#9c36b5", "#0c8599", "#5c940d", "#c2255c", "#495057"]
    for x, y, t, hh in bt.texts:
        if can_poly.buffer(1500).contains(Point(x, y)):
            ax.text(x, y, t, fontsize=5, color="#5f3dc4")
    for p1, p2, p3 in bt.dims:
        if not can_poly.buffer(1500).contains(Point(p1)):
            continue
        vx, vy = p3[0] - p1[0], p3[1] - p1[1]
        q1, q2 = (p1[0] + vx, p1[1] + vy), (p2[0] + vx, p2[1] + vy)
        ax.plot([q1[0], q2[0]], [q1[1], q2[1]], color="#868e96", lw=0.4)
        ax.text((q1[0] + q2[0]) / 2, (q1[1] + q2[1]) / 2, f"{math.dist(p1, p2):.0f}", fontsize=4.5, color="#495057", ha="center")
    for d in bt.ds:
        if d["can"] != ten_can:
            continue
        c = d["cat"]
        rw, sau = c.get("rong", 200), c.get("sau", 160)
        nx, ny = d["n"]
        ux, uy = ny, -nx
        q = [(d["x"] + ux * rw / 2, d["y"] + uy * rw / 2), (d["x"] - ux * rw / 2, d["y"] - uy * rw / 2),
             (d["x"] - ux * rw / 2 + nx * sau, d["y"] - uy * rw / 2 + ny * sau), (d["x"] + ux * rw / 2 + nx * sau, d["y"] + uy * rw / 2 + ny * sau)]
        col = "#e03131" if d.get("btct") else "#0b7285"
        ax.fill([p[0] for p in q], [p[1] for p in q], color=col, alpha=0.85)
        lab = (c.get("gia_tri") or c["ma"].split("-")[-1])
        ax.text(d["x"] + nx * (sau + 90), d["y"] + ny * (sau + 90), lab, fontsize=5.5, color="#0b7285", ha="center", va="center")
    for x, y in bt.may:
        if can_poly.buffer(300).contains(Point(x, y)):
            ax.add_patch(matplotlib.patches.Circle((x, y), 250, fill=False, ec="#e03131", lw=1, ls="--"))
    if co_san:
        for d in co_san:
            if can_poly.buffer(300).contains(Point(d["x"], d["y"])):
                ax.plot(d["x"], d["y"], "s", mfc="none", mec="#e03131" if d.get("loi") else "#2b8a3e", ms=6, mew=1.2)
    if de_xuat:
        for x, y, t in de_xuat:
            if can_poly.buffer(300).contains(Point(x, y)):
                ax.plot(x, y, "o", mfc="none", mec="#1c7ed6", ms=9, mew=1.5)
    ax.set_aspect("equal")
    ax.set_xlim(minx - 800, maxx + 800)
    ax.set_ylim(miny - 800, maxy + 800)
    ax.axis("off")
    ax.set_title(f"{ten_can} – xanh đậm: thiết bị (đỏ: trên vách BTCT, mây = đặt chờ) | nét chấm gạch: dây theo lộ | "
                 "hồng: vùng cấm (cửa, sau cánh cửa, sau tủ áo) | xám: vách BTCT | cam #n: block chưa nhận diện", fontsize=7)
    fig.savefig(path)
    plt.close(fig)


def xuat_excel(path, du_an, file, bt, so, rooms_all, nt, chua_ro, che_do, co_san=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Tom tat"
    mas = sorted({d["ma"] for d in bt.ds} | ({d["ma"] for d in co_san} if co_san else set()))
    ws.append(["Căn hộ", "Phòng", "Loại"] + mas + ["Tổng", "Lỗi", "Cảnh báo", "Gợi ý", "Cần hỏi"])
    for r in rooms_all:
        c = Counter(d["ma"] for d in (co_san if che_do == "soat" else bt.ds) if d["can"] == r["can"] and d["phong"] == r["ten"])
        m = Counter(x["muc"] for x in so.ds if x["can"] == r["can"] and x["phong"] == r["ten"])
        hoi = sum(1 for x in so.ds if x["can"] == r["can"] and x["phong"] == r["ten"] and x["hoi"])
        ws.append([r["can"], r["ten"], ", ".join(r["loai"])] + [c.get(k, 0) or None for k in mas] + [sum(c.values()), m[LOI], m[CB], m[GY], hoi])
    ws2 = wb.create_sheet("Loi va can hoi")
    ws2.append(["STT", "Dự án", "File/Tầng", "Căn hộ", "Phòng", "Hạng mục", "Mô tả", "Thiết bị", "X", "Y", "Đề xuất X", "Đề xuất Y",
                "Giá trị đo (mm)", "Ngưỡng (mm)", "Mức độ", "Cần hỏi người dùng", "Người phụ trách", "Trạng thái"])
    for i, x in enumerate(so.ds, 1):
        d = x["tb"]
        ws2.append([i, du_an, file, x["can"], x["phong"], x["hang_muc"], x["mo_ta"], d["cat"]["ten"] if d else "",
                    round(d["x"]) if d else None, round(d["y"]) if d else None,
                    round(x["de_xuat"][0]) if x["de_xuat"] else None, round(x["de_xuat"][1]) if x["de_xuat"] else None,
                    x["gia_tri"], x["nguong"], x["muc"], "Có" if x["hoi"] else "", "", "CẦN HỎI" if x["hoi"] else "CẦN XEM"])
    mau = {LOI: "FCE4D6", CB: "FFF2CC", GY: "E2EFDA"}
    for row in ws2.iter_rows(min_row=2):
        row[14].fill = PatternFill("solid", fgColor=mau.get(row[14].value, "FFFFFF"))
    ws3 = wb.create_sheet("Thiet bi " + ("de xuat" if che_do == "soat" else "bo tri"))
    ws3.append(["STT", "Căn hộ", "Phòng", "Mã", "Thiết bị", "Ký hiệu", "Block", "Layer", "Cao độ", "X", "Y", "Xoay", "Lộ", "Nguyên tắc / lý do", "Ghi chú"])
    for i, d in enumerate(sorted(bt.ds, key=lambda d: (d["can"], d["phong"], d["ma"])), 1):
        c = d["cat"]
        ws3.append([i, d["can"], d["phong"], d["ma"], c["ten"], c.get("gia_tri") or "", c["block"], bt.cfg["ve"]["layer"][c["layer"]][0], c.get("cao_do", ""), round(d["x"]), round(d["y"]),
            round(d["rot"]), d.get("lo") or "", d["ly_do"], d.get("ghi_chu") or ""])
    if co_san is not None:
        ws5 = wb.create_sheet("Thiet bi hien co")
        ws5.append(["STT", "Căn hộ", "Phòng", "Mã", "Ký hiệu", "Block", "Layer", "Handle", "X", "Y", "Xoay", "Lỗi / ghi chú"])
        for i, d in enumerate(co_san, 1):
            ws5.append([i, d.get("can", ""), d.get("phong", ""), d["ma"], d.get("val") or "", d["block"], d["layer"], d["handle"],
                        round(d["x"]), round(d["y"]), round(d["rot"]), "; ".join(d.get("loi_ds", []))])
    ws4 = wb.create_sheet("Lo")
    ws4.append(["Căn hộ", "Lộ", "Nhãn", "Số thiết bị", "Thành phần", "Ghi chú"])
    lo = defaultdict(list)
    for d in bt.ds:
        if d.get("lo"):
            lo[(d["can"], d["lo"])].append(d)
    for (can, l), mem in sorted(lo.items()):
        c = Counter(d["cat"].get("gia_tri") or d["ma"] for d in mem)
        ws4.append([can, l, f"{l}/TĐ.CH", len(mem), ", ".join(f"{k}×{v}" for k, v in c.items()),
                    "Lộ riêng" if re.match(r"^(AC|BT|HW)", l) else "Kỹ sư điện kiểm tra tải theo CB và tiết diện dây"])
    ws6 = wb.create_sheet("Noi that")
    ws6.append(["Loại", "Tên block", "X", "Y", "Rộng", "Sâu", "Ghi chú"])
    for f in nt:
        b_ = f["fp"].bounds
        ws6.append([f["loai"], f.get("ten", ""), round(f["x"]), round(f["y"]), round(b_[2] - b_[0]), round(b_[3] - b_[1]),
                    "suy theo hình dạng – xác nhận" if f.get("suy") else ""])
    for k in chua_ro:
        ws6.append([f"CHƯA NHẬN DIỆN #{k['id']}", k["ten"], round(k["x"]), round(k["y"]), round(k["rong"]), round(k["cao"]),
                    "Gán loại trong file --nhan-dien-bo-sung nếu là thiết bị cần ổ cắm"])
    for w in wb.worksheets:
        for c in w[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for col in w.columns:
            L = max(len(str(c.value or "")) for c in col)
            w.column_dimensions[col[0].column_letter].width = min(60, max(7, L * 0.9))
    wb.save(path)


# ================================================================================================= soat
def soat_co_san(bt, co_san, rooms, cuas, nt, dims, nhan_lo, so):
    """So thiet bi da ve voi phuong an theo nguyen tac (bt.ds) + kiem tra rang buoc cung."""
    b = bt.b
    for d in co_san:
        d["loi_ds"] = []
        p = Point(d["x"], d["y"])
        r = next((r for r in rooms if r["poly"].buffer(80).contains(p)), None)
        if r is None:
            d["can"], d["phong"] = "?", "(ngoài phòng)"
            continue
        d["can"], d["phong"] = r["can"], r["ten"]
        bien = r["bien"]
        s0 = bien.chieu((d["x"], d["y"]))
        kc_t = bien.ring.distance(p)
        if d["ma"] not in ("VDP",) and kc_t > 80 and not d.get("lech_quy_uoc"):
            so.them(r["can"], r["ten"], CB, "Thiết bị không bám tường", f"{d['cat']['ten']} cách mặt tường {kc_t:.0f} mm.", tb=d, gia_tri=round(kc_t), nguong=80)
            d["loi_ds"].append("không bám tường")
        ld = bien.ly_do_cam(s0)
        # chi tiet lap dat (08/10/2026): cong tac cach lo mo toi thieu 80 (khong ap nguong 200 cua o cam); tu dien dat phia
        # ban le sat goc (nguoi dung chon) -> khong xet 'sau canh cua' / 'cach khuon < 200'
        mien = {"CT-20A": ("cách khuôn cửa < 200",), "CT-BA": ("cách khuôn cửa < 200",),
                "TU-DIEN": ("cách khuôn cửa < 200", "sau cánh cửa")}.get(d["ma"], ())
        if ld and ld not in mien and not d.get("lech_quy_uoc") and d["ma"] not in ("VDP", "MUI-TEN"):
            so.them(r["can"], r["ten"], LOI, "Vị trí vi phạm", f"{d['cat']['ten']} ({d.get('val') or ''}): {ld}.", tb=d)
            d["loi_ds"].append(ld)
            d["loi"] = True
        if bien.la_btct(s0) and not any(math.dist((d["x"], d["y"]), m) < 400 for m in bt.may_co_san):
            so.them(r["can"], r["ten"], CB, "Thiết bị trên vách BTCT", f"{d['cat']['ten']} nằm trên vách BTCT, chưa đánh dấu 'đặt chờ khi đổ cột vách'. Ưu tiên dời sang tường xây.", tb=d)
            d["loi_ds"].append("vách BTCT")
        if d["ma"] not in ("VDP", "MUI-TEN", "QUAT-HUT") and not d.get("lech_quy_uoc") and not any(math.dist((d["x"], d["y"]), q) < 30 for q in dims):
            so.them(r["can"], r["ten"], GY, "Thiếu kích thước định vị", f"{d['cat']['ten']} chưa có dim từ mép tường tới tâm thiết bị.", tb=d)
    # doi chieu voi phuong an theo nguyen tac
    de_xuat = []
    dung = set()
    for e in bt.ds:
        if e["ma"] in ("MUI-TEN",):
            continue
        if e["phong"] == "Ngoài cửa chính":
            cands = [d for d in co_san if d["ma"] == "VDP" and id(d) not in dung and math.dist((d["x"], d["y"]), (e["x"], e["y"])) < 2500]
        else:
            cands = [d for d in co_san if d["ma"] == e["ma"] and d.get("can") == e["can"] and d.get("phong") == e["phong"] and id(d) not in dung]
        if not cands:
            muc = LOI if e["ma"] in ("O-G", "O-TV", "O-TL", "HOP-BT", "HOP-HM", "HOP-BNL", "O-W", "TU-DIEN", "HOP-AC", "O-B", "CT-20A") else CB
            so.them(e["can"], e["phong"], muc, "Thiếu thiết bị", f"Thiếu {e['cat']['ten']} ({e['ly_do']}).", de_xuat=(e["x"], e["y"]))
            de_xuat.append((e["x"], e["y"], e["ma"]))
            continue
        d = min(cands, key=lambda d: math.dist((d["x"], d["y"]), (e["x"], e["y"])))
        dung.add(id(d))
        kc = math.dist((d["x"], d["y"]), (e["x"], e["y"]))
        if kc > 300:
            so.them(e["can"], e["phong"], CB, "Lệch so với nguyên tắc", f"{e['cat']['ten']} lệch {kc:.0f} mm so với vị trí theo nguyên tắc ({e['ly_do']}).",
                    tb=d, de_xuat=(e["x"], e["y"]), gia_tri=round(kc), nguong=300)
            de_xuat.append((e["x"], e["y"], e["ma"]))
    for d in co_san:
        if id(d) not in dung and d.get("phong") not in (None, "(ngoài phòng)") and d["ma"] not in ("MUI-TEN", "MAY-CHO", "QUAT-HUT"):
            so.them(d.get("can"), d.get("phong"), GY, "Thiết bị ngoài nguyên tắc", f"{d['cat']['ten']} ({d.get('val') or ''}) không có trong phương án theo nguyên tắc: kiểm tra có cần không.", tb=d)
    # (08/10/2026) khong ve / khong soat day, nhan lo: lo chi ghi Excel
    return de_xuat


# ================================================================================================= main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("che_do", choices=["bo-tri", "soat"])
    ap.add_argument("dxf")
    ap.add_argument("--out-dir", default=".")
    ap.add_argument("--du-an", default="")
    ap.add_argument("--cau-hinh", default=os.path.join(HERE, "cau_hinh_o_cam.json"))
    ap.add_argument("--catalog", default=os.path.join(THU_VIEN, "catalog.json"))
    ap.add_argument("--so-o-bep", default=None, help='so o cam mat bep: "2" hoac "CH01=2,CH02=3" (HOI NGUOI DUNG)')
    ap.add_argument("--de-o", choices=["chu_nhat", "vuong"], default="vuong",
                    help="de o cam: TV-DN tam cach 100 (vuong, hop 86x86 - chi tiet lap dat) / 150 (chu nhat)")
    ap.add_argument("--can", default="", help="chi xu ly cac can (vd CH01,CH02) - can dien hinh")
    ap.add_argument("--may-rua-bat", action="store_true")
    ap.add_argument("--lo-nuong", action="store_true")
    ap.add_argument("--tv-pn-theo-truc-giuong", action="store_true")
    ap.add_argument("--nhan-dien-bo-sung", default="", help='JSON {"ten block": "loai noi that"}')
    ap.add_argument("--layer-ten", default="", help="layer Text ten phong (mac dinh theo cau hinh)")
    ap.add_argument("--khong-bang-ky-hieu", action="store_true")
    ap.add_argument("--nen-revit", choices=["auto", "co", "khong"], default="auto",
                    help="nen xuat tu Revit (net no, khong ten phong): dung phong / can / noi that theo hinh dang (nen_revit.py)")
    ap.add_argument("--bo-can", default="", help="bo qua cac can (vd can da ve san)")
    ap.add_argument("--bnl-mac-dinh", action="store_true",
                    help="WC khong co block binh nong lanh: dat box BNL o goc tuong phia ban le cua WC (theo can mau nguoi dung ve)")
    a = ap.parse_args()
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    cfg = json.load(open(a.cau_hinh, encoding="utf-8"))
    catalog = json.load(open(a.catalog, encoding="utf-8"))
    if a.layer_ten:
        cfg["layer_ten_phong"] = a.layer_ten
    bo_sung = ap_nhan_dien_bo_sung(cfg, a.nhan_dien_bo_sung)
    tpp._BO_QUA[0] = tpp.BLOCK_BO_QUA
    t0 = time.time()
    doc = ezdxf.readfile(a.dxf)
    msp = doc.modelspace()
    so = SoTay()
    van_de = []
    if doc.header.get("$INSUNITS", 0) != 4:
        van_de.append(f"INSUNITS={doc.header.get('$INSUNITS', 0)}, chuẩn Archivina là 4 (mm).")
    ca = box(-1e9, -1e9, 1e9, 1e9)
    nt = st.doc_noi_that(msp, doc, cfg, ca)
    print(f"[noi that] {len(nt)} {time.time() - t0:.0f}s", file=sys.stderr)
    nguon = msp            # nguon doi tuong cho doc_cua / doc_hinh_lop (nen Revit: danh sach da trai phang, doc mot lan)
    ds_phong, cans = ([], []) if a.nen_revit == "co" else st.dung_phong(msp, doc, cfg, [], nt)
    tk_revit = None
    if not cans and a.nen_revit != "khong":
        import nen_revit as nr
        rv = cfg["nen_revit"]
        nguon = list(tpp.walk(msp))
        net, texts = nr.doc_net(nguon, tpp, rv["layer_tuong"] + rv["layer_cua"] + rv["layer_noi_that"] + rv["layer_thiet_bi_ve_sinh"]
                                + rv["layer_dan_nong"] + rv["layer_ma_can"] + rv.get("layer_cung_cua_them", []), rv["layer_ma_can"])
        cfg["layer_cung_cua_them"] = cfg.get("layer_cung_cua_them", []) + rv.get("layer_cung_cua_them", [])
        print(f"[revit] doc net {time.time() - t0:.0f}s", file=sys.stderr)
        nt_rv, tk_revit = nr.nhan_dien_noi_that(net, rv)
        ma = nr.doc_ma_can(net, texts, rv)
        ds_phong, cans, gc = nr.dung_phong_can(net, nt_rv, ma, rv, tpp, log=lambda s: print(s, file=sys.stderr))
        nt = nt + [f for f in nt_rv if not f.get("chi_phong_ngu") or
                   any("ngu" in r["loai"] and r["poly"].buffer(100).contains(Point(f["x"], f["y"])) for r in ds_phong)]
        van_de.append(f"Nền xuất từ Revit (nét đã nổ, không có Text tên phòng): {len(cans)} căn theo mã căn + đường dẫn, phòng dựng từ "
                      f"nét tường / cửa, loại phòng và nội thất nhận theo hình dạng ({tk_revit['loai']}) – kiểm tra ảnh từng căn.")
        van_de += gc
        print(f"[revit] {len(ds_phong)} phong, {len(cans)} can {time.time() - t0:.0f}s", file=sys.stderr)
    print(f"[phong] {len(ds_phong)} {time.time() - t0:.0f}s", file=sys.stderr)
    if not cans:
        raise SystemExit("Không dựng được phòng/căn nào: kiểm tra layer tên phòng, nét tường (cau_hinh_o_cam.json).")
    loc = {bo_dau(x) for x in a.can.split(",") if x.strip()}
    if loc:
        cans = [(c, t) for c, t in cans if bo_dau(t) in loc]
    bo = {bo_dau(x) for x in a.bo_can.split(",") if x.strip()}
    if bo:
        cans = [(c, t) for c, t in cans if bo_dau(t) not in bo]
    rooms = [r for r in ds_phong if r["poly"] is not None and any(t == r["can"] for _, t in cans)]
    for r in ds_phong:
        if r["poly"] is None and any(t == r["can"] for _, t in cans):
            so.them(r["can"], r["ten"], CB, "Không dựng được ranh phòng", "Nét tường/vách hở lớn: phòng chưa bố trí được.", hoi=True)
    if any(r.get("suy_ra") for r in rooms):
        van_de.append(f"{sum(1 for r in rooms if r.get('suy_ra'))} phòng không có Text tên phòng: loại phòng suy theo nội thất – kiểm tra lại.")
        for r in rooms:
            if r.get("nen_revit"):
                if r["loai"] == ["khac"]:
                    so.them(r["can"], r["ten"], GY, "Phòng chưa xác định loại", f"Không nhận được nội thất đặc trưng: chưa bố trí ổ cắm cho {r['ten']} ({r['poly'].area / 1e6:.1f} m²).")
            elif r.get("suy_ra"):
                so.them(r["can"], r["ten"], CB, "Thiếu tên phòng", f"Không có Text tên phòng: loại '{r['ten']}' suy theo nội thất.", hoi=not r["loai"] or r["loai"] == ["khac"])
    vung = unary_union([c for c, _ in cans] + [r["poly"] for r in rooms]).buffer(500)
    nt = [f for f in nt if vung.intersects(f["fp"])]
    # bien phong + vung cam
    cuas = doc_cua(nguon, cfg, rooms)
    net_cua = unary_union(doc_hinh_lop(nguon, cfg["layer_cua"]))
    net_btct = unary_union(doc_hinh_lop(nguon, cfg["layer_btct"]))
    net_lc = unary_union(doc_hinh_lop(nguon, cfg.get("layer_lan_can", [])))
    net_tuong = doc_hinh_lop(nguon, cfg["layer_ranh"])
    for r in rooms:
        bien = Bien(r["poly"])
        r["poly"] = bien.P
        r["bien"] = bien
        P = r["poly"]
        Pb = P.buffer(300)
        bien.danh_dau(net_cua.intersection(Pb) if not net_cua.is_empty else None, 100, "cửa / cửa sổ")
        bien.danh_dau(net_lc.intersection(Pb) if not net_lc.is_empty else None, 80, "lan can / vách kính")
        for c in cuas:
            if c["mo"].distance(P.exterior) < 300:
                bien.danh_dau(c["mo"], cfg["bo_tri"]["o_cach_khuon_cua_toi_thieu"], "cách khuôn cửa < 200")
            if P.contains(Point(c["mid"])):
                bien.danh_dau(Point(c["H"]), c["R"] + 50, "sau cánh cửa")
        for f in nt:
            if f["loai"] == "tu_ao" and P.buffer(200).intersects(f["fp"]) and P.buffer(150).contains(f["fp"].representative_point()):
                bien.danh_dau(f["fp"], 250, "sau tủ áo")
        if not net_btct.is_empty:
            bien.danh_dau_btct(net_btct.intersection(Pb))
    # ke TV suy theo hinh dang + block chua nhan dien
    chua_ro = doc_khoi_chua_ro(msp, doc, cfg, vung, nt)
    kt = cfg["ke_tv_suy_theo_hinh"]
    ke_tv_suy = []
    for k in chua_ro:
        dai, sau = max(k["rong"], k["cao"]), min(k["rong"], k["cao"])
        if not (kt["dai"][0] <= dai <= kt["dai"][1] and kt["sau"][0] <= sau <= kt["sau"][1]):
            continue
        r = next((r for r in rooms if r["poly"].contains(Point(k["x"], k["y"])) and {"ngu", "khach"} & set(r["loai"])), None)
        if r is None or r["poly"].exterior.distance(k["fp"]) > kt["sat_tuong"]:
            continue
        if any(f["loai"] == "ke_tv" and r["poly"].contains(Point(f["x"], f["y"])) for f in nt):
            continue
        ref = [f for f in nt if f["loai"] in ("giuong", "sofa") and r["poly"].contains(Point(f["x"], f["y"]))]
        if ref:
            # tuong cua ke phai doi dien giuong / sofa: phap tuyen huong ve phia noi that
            _, kk, _, _, _ = tuong_gan(r["bien"], k["fp"])
            v = (ref[0]["x"] - k["x"], ref[0]["y"] - k["y"])
            if v[0] * kk["n"][0] + v[1] * kk["n"][1] < 0.5 * math.hypot(*v):
                continue
        k["suy"] = True
        ke_tv_suy.append(dict(loai="ke_tv", ten=k["ten"], fp=k["fp"], x=k["x"], y=k["y"], suy=True))
    ten_suy = {f["ten"] for f in ke_tv_suy}
    chua_ro = [k for k in chua_ro if not k.get("suy")]
    for i, k in enumerate(chua_ro, 1):
        k["id"] = i
    nt_all = nt + ke_tv_suy
    for f in ke_tv_suy:
        rr = next((r for r in rooms if r["poly"].contains(Point(f["x"], f["y"]))), None)
        so.them(rr["can"] if rr else "", rr["ten"] if rr else "", GY, "Kệ TV suy theo hình dạng",
                f"Block '{f['ten']}' ({f['fp'].bounds[2] - f['fp'].bounds[0]:.0f}×{f['fp'].bounds[3] - f['fp'].bounds[1]:.0f}) mỏng, sát tường đối diện giường/sofa: coi là kệ TV – xác nhận.")
    # bo tri theo nguyen tac
    bt = BoTri(cfg, catalog, a, so)
    lst = list(net_tuong)
    bt.tuong_net = (lst, shapely.STRtree(lst)) if lst else None
    bt.quat = []
    bt.may_co_san = []
    so_pn = {}
    for poly, ten in cans:
        rr = [r for r in rooms if r["can"] == ten]
        if not rr:
            continue
        so_pn[ten] = bo_tri_can(bt, ten, rr, nt_all, cuas, poly, a, ke_tv_suy)
    # thiet bi dien co san (quat hut de noi F; soat)
    co_san = doc_o_cam_co_san(msp, doc, catalog, vung)
    for d in co_san:
        r = next((r for r in rooms if r["poly"].buffer(80).contains(Point(d["x"], d["y"]))), None)
        d["can"], d["phong"] = (r["can"], r["ten"]) if r else ("?", "(ngoài phòng)")
    bt.quat = [d for d in co_san if d["ma"] == "QUAT-HUT"] + [
        dict(x=f["x"], y=f["y"], can=next((r["can"] for r in rooms if r["poly"].contains(Point(f["x"], f["y"]))), "?"),
             phong=next((r["ten"] for r in rooms if r["poly"].contains(Point(f["x"], f["y"]))), "?")) for f in nt if f["loai"] == "quat_hut"]
    bt.may_co_san = [(d["x"], d["y"]) for d in co_san if d["ma"] == "MAY-CHO"]
    for poly, ten in cans:
        rr = [r for r in rooms if r["can"] == ten]
        if rr:
            tach = chia_lo(bt, ten, rr, so_pn.get(ten, 0))
            if tach:
                so.them(ten, "", GY, "Tách 2 lộ phòng ngủ", f"Căn có {so_pn.get(ten)} phòng ngủ (> 3): ổ cắm phòng ngủ tách 2 lộ.")
            ve_dim(bt, ten)
    os.makedirs(a.out_dir, exist_ok=True)
    tien_to = "BoTri" if a.che_do == "bo-tri" else "Soat"
    de_xuat = None
    if a.che_do == "soat":
        dims, nhan_lo = doc_dim_va_nhan(msp)
        dt = [d for d in co_san if d["ma"] not in ("MUI-TEN", "MAY-CHO")]
        for d in dt:
            # box BT ve lech 200 vao phong truoc box HM (quy uoc ban ve mau): khong xet bam tuong / dim rieng
            # ky hieu ve lech vao phong theo quy uoc: BT truoc HM, o cam H400 duoi man hinh chuong cua
            d["lech_quy_uoc"] = (d["ma"] == "HOP-BT" and any(x["ma"] == "HOP-HM" and math.dist((x["x"], x["y"]), (d["x"], d["y"])) < 300 for x in dt)) or                 (d["ma"] == "O-DOI" and any(x["ma"] == "VDP" and math.dist((x["x"], x["y"]), (d["x"], d["y"])) < 350 for x in dt))
        if not dt:
            van_de.append("Không tìm thấy ký hiệu ổ cắm / hộp chờ nào theo thư viện (catalog.json → bi_danh_block): kiểm tra tên block.")
        # che do soat: ghi chu cua PHUONG AN THEO NGUYEN TAC (BTCT, dich vi tri...) khong phai loi cua ban ve -> bo;
        # giu cac muc thieu du lieu / can hoi (phan chua kiem tra duoc)
        so.ds = [x for x in so.ds if x["hoi"] or x["hang_muc"] in ("Kệ TV suy theo hình dạng", "Thiếu tên phòng")]
        for x in so.ds:
            if x["hoi"]:
                x["mo_ta"] = "Chưa kiểm tra được đầy đủ: " + x["mo_ta"]
        de_xuat = soat_co_san(bt, dt, rooms, cuas, nt_all, dims, nhan_lo, so)
    xlsx = os.path.join(a.out_dir, f"BaoCao{tien_to}OCam.xlsx")
    xuat_excel(xlsx, a.du_an, os.path.basename(a.dxf), bt, so, rooms, nt_all, chua_ro, a.che_do, co_san if a.che_do == "soat" else None)
    anh = []
    for poly, ten in cans:
        rr = [r for r in rooms if r["can"] == ten]
        if not rr:
            continue
        p = os.path.join(a.out_dir, f"xem_o_cam_{bo_dau(ten).replace(' ', '_')}.png")
        ve_anh(p, poly, rr, nt_all, bt, cuas, so, ten, chua_ro, co_san if a.che_do == "soat" else None, de_xuat)
        anh.append(p)
    scr = None
    if a.che_do == "bo-tri":
        scr = os.path.join(a.out_dir, "ve_o_cam.scr")
        mx = max(c.bounds[2] for c, _ in cans)
        my = max(c.bounds[3] for c, _ in cans)
        xuat_scr(scr, bt, cfg, catalog, not a.khong_bang_ky_hieu, (mx + 2000, my - 4600))
    js = dict(
        che_do=a.che_do, file=os.path.basename(a.dxf), van_de=van_de, nhan_dien_bo_sung=bo_sung,
        tham_so=dict(so_o_bep=a.so_o_bep, de_o=a.de_o, may_rua_bat=a.may_rua_bat, lo_nuong=a.lo_nuong,
                     tv_pn_theo_truc_giuong=a.tv_pn_theo_truc_giuong),
        can_ho=[dict(ten=t, so_phong=sum(1 for r in rooms if r["can"] == t), so_phong_ngu=so_pn.get(t, 0),
                     thiet_bi=dict(Counter(d["ma"] for d in bt.ds if d["can"] == t)),
                     lo=sorted({d["lo"] for d in bt.ds if d["can"] == t and d.get("lo")})) for _, t in cans],
        phong=[dict(can=r["can"], ten=r["ten"], loai=r["loai"], dt=round(r["poly"].area / 1e6, 1)) for r in rooms],
        noi_that=dict(Counter(f["loai"] for f in nt_all)),
        block_chua_nhan_dien=[dict(id=k["id"], ten=k["ten"], x=round(k["x"]), y=round(k["y"]), rong=round(k["rong"]), sau=round(k["cao"])) for k in chua_ro],
        can_hoi=[dict(can=x["can"], phong=x["phong"], hang_muc=x["hang_muc"], mo_ta=x["mo_ta"]) for x in so.ds if x["hoi"]],
        muc_do=dict(Counter(x["muc"] for x in so.ds)),
        thiet_bi_co_san=dict(Counter(d["ma"] for d in co_san)) if co_san else {},
        xlsx=xlsx, anh=anh, scr=scr, thoi_gian_s=round(time.time() - t0))
    js_path = os.path.join(a.out_dir, "bo_tri_o_cam.json" if a.che_do == "bo-tri" else "soat_o_cam.json")
    full = dict(js, thiet_bi=[dict(ma=d["ma"], block=d["cat"]["block"], gia_tri=d["cat"].get("gia_tri"), x=round(d["x"], 1), y=round(d["y"], 1),
                                   rot=round(d["rot"], 2), can=d["can"], phong=d["phong"], lo=d.get("lo"), ly_do=d["ly_do"]) for d in bt.ds],
                dims=[[list(map(lambda v: round(v, 1), p)) for p in q] for q in bt.dims])
    json.dump(full, open(js_path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    js["json"] = js_path
    print(json.dumps(js, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
